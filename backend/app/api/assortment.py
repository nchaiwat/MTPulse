import json
from decimal import Decimal
from io import BytesIO
from typing import Annotated, Literal
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.auth import actor_from_session
from app.database import get_session
from app.models import (
    AssortmentBase,
    AssortmentForecast,
    AssortmentMember,
    AssortmentPlan,
    AuthUser,
    ItemMapping,
    TransactionLog,
)
from app.modern_trade_registry import ActiveReportingModernTradeCode
from app.services import assortment
from app.services.audit_request import request_details
from app.services.telegram import set_setting

router = APIRouter(prefix="/api/assortment", tags=["Assortment"])
Db = Annotated[Session, Depends(get_session)]


def audit(db, request, action, detail):
    actor = actor_from_session(db, "development-admin")
    if actor.startswith("user:"):
        user = db.get(AuthUser, actor[5:])
        if user:
            actor = "user:" + user.username
    db.add(
        TransactionLog(
            event_code="ASRT-01",
            category="assortment",
            action=action,
            status="success",
            message="Assortment: " + action,
            triggered_by=actor,
            details=json.dumps({**request_details(request), **detail}, ensure_ascii=False),
            records_count=1,
            duration_ms=0,
        )
    )


def commit(db):
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "ข้อมูลเปลี่ยนแล้ว กรุณาโหลดใหม่") from exc


class BaseInput(BaseModel):
    description: str = Field(min_length=1, max_length=500)
    attributes: list[str] = Field(default_factory=lambda: [""] * 7, min_length=7, max_length=7)
    members: list[str] = Field(default_factory=list, max_length=500)
    version: int = Field(default=0, ge=0)

    @field_validator("description")
    @classmethod
    def description_valid(cls, v):
        if not v.strip():
            raise ValueError("ต้องระบุ Description")
        return v.strip()

    @field_validator("attributes", "members")
    @classmethod
    def fields_valid(cls, values):
        if any(len(v) > 200 for v in values):
            raise ValueError("ข้อความยาวเกินกำหนด")
        return [v.strip() for v in values]


@router.get("/catalog")
def catalog(db: Db):
    return assortment.catalog(db)


@router.get("/suggestions")
def suggestions(
    db: Db, search: str = Query(default="", max_length=200), page: int = Query(default=1, ge=1)
):
    return assortment.candidates(db, search, page)


def save_base(db, request, payload, row=None):
    if len(set(payload.members)) != len(payload.members):
        raise HTTPException(422, "WA Item ซ้ำ")
    codes = set(
        db.scalars(
            select(ItemMapping.wa_item_code).where(
                ItemMapping.wa_item_code.in_(payload.members), ItemMapping.status == "confirmed"
            )
        )
    )
    if set(payload.members) != codes:
        raise HTTPException(422, "ไม่พบ WA Item ที่ยืนยันแล้ว")
    if row and row.version != payload.version:
        raise HTTPException(409, "Base Item ถูกแก้ไขแล้ว กรุณาโหลดใหม่")
    if row is None:
        row = AssortmentBase(id=str(uuid4()), version=1)
    before = {
        "description": row.description,
        "members": list(
            db.scalars(
                select(AssortmentMember.wa_item_code).where(AssortmentMember.base_id == row.id)
            )
        ),
    }
    conflicts = list(
        db.scalars(
            select(AssortmentMember).where(
                AssortmentMember.wa_item_code.in_(payload.members),
                AssortmentMember.base_id != row.id,
            )
        )
    )
    if conflicts:
        raise HTTPException(409, "WA Item อยู่ใน Base Item อื่นแล้ว")
    if row in db:
        row.version += 1
    db.add(row)
    row.description, row.attributes = payload.description, json.dumps(payload.attributes)
    db.flush()
    db.execute(delete(AssortmentMember).where(AssortmentMember.base_id == row.id))
    db.add_all([AssortmentMember(wa_item_code=wa, base_id=row.id) for wa in payload.members])
    audit(
        db,
        request,
        "confirm_base",
        {"base_id": row.id, "before": before, "after": payload.model_dump()},
    )
    commit(db)
    return {"id": row.id, "version": row.version}


@router.post("/bases")
def create_base(payload: BaseInput, request: Request, db: Db):
    return save_base(db, request, payload)


@router.put("/bases/{base_id}")
def update_base(base_id: str, payload: BaseInput, request: Request, db: Db):
    row = db.scalar(select(AssortmentBase).where(AssortmentBase.id == base_id).with_for_update())
    if row is None:
        raise HTTPException(404, "ไม่พบ Base Item")
    return save_base(db, request, payload, row)


class PlanInput(BaseModel):
    name: str = Field(min_length=1, max_length=150)
    year: int = Field(ge=2002, le=2100)
    copy_id: str | None = None


@router.post("/plans")
def create_plan(payload: PlanInput, request: Request, db: Db):
    if not payload.name.strip():
        raise HTTPException(422, "ต้องระบุชื่อแผน")
    plan = AssortmentPlan(id=str(uuid4()), name=payload.name.strip(), year=payload.year, version=1)
    db.add(plan)
    if payload.copy_id:
        source = db.scalar(
            select(AssortmentPlan).where(AssortmentPlan.id == payload.copy_id).with_for_update()
        )
        if not source or source.year != payload.year:
            raise HTTPException(422, "คัดลอกได้เฉพาะแผนปีเดียวกัน")
        db.flush()
        for f in db.scalars(
            select(AssortmentForecast).where(AssortmentForecast.plan_id == source.id)
        ):
            db.add(
                AssortmentForecast(
                    plan_id=plan.id, base_id=f.base_id, mt_code=f.mt_code, months=f.months
                )
            )
    audit(db, request, "create_plan", {"plan_id": plan.id, **payload.model_dump()})
    commit(db)
    return {"id": plan.id, "version": plan.version}


class VersionInput(BaseModel):
    version: int = Field(ge=1)


def lock_plan(db, plan_id, version):
    plan = db.scalar(select(AssortmentPlan).where(AssortmentPlan.id == plan_id).with_for_update())
    if not plan:
        raise HTTPException(404, "ไม่พบแผน")
    if plan.version != version:
        raise HTTPException(409, "แผนถูกแก้ไขแล้ว กรุณาโหลดใหม่ก่อนบันทึก")
    return plan


@router.put("/plans/{plan_id}/primary")
def primary(plan_id: str, payload: VersionInput, request: Request, db: Db):
    # Stable ordering prevents competing primary selections from locking in reverse order.
    list(db.scalars(select(AssortmentPlan).order_by(AssortmentPlan.id).with_for_update()))
    plan = lock_plan(db, plan_id, payload.version)
    previous = db.scalar(select(AssortmentPlan).where(AssortmentPlan.primary_year == plan.year))
    if previous:
        previous.primary_year = None
        previous.version += 1
        db.flush()
    plan.primary_year = plan.year
    plan.version += 1
    audit(db, request, "primary_plan", {"plan_id": plan.id, "year": plan.year})
    commit(db)
    return {"version": plan.version}


class ForecastInput(VersionInput):
    months: list[Decimal | None] = Field(min_length=12, max_length=12)

    @field_validator("months")
    @classmethod
    def valid_months(cls, values):
        if any(
            v is not None
            and (
                not v.is_finite()
                or v < 0
                or v > Decimal("1000000000")
                or v.as_tuple().exponent < -4
            )
            for v in values
        ):
            raise ValueError("จำนวนต้องเป็น 0–1,000,000,000 ทศนิยมไม่เกิน 4 ตำแหน่ง")
        return values


@router.put("/plans/{plan_id}/forecast/{base_id}/{mt}")
def save_forecast(
    plan_id: str,
    base_id: str,
    mt: ActiveReportingModernTradeCode,
    payload: ForecastInput,
    request: Request,
    db: Db,
):
    plan = lock_plan(db, plan_id, payload.version)
    if not db.get(AssortmentBase, base_id):
        raise HTTPException(404, "ไม่พบ Base Item")
    row = db.get(AssortmentForecast, (plan_id, base_id, mt))
    before = json.loads(row.months) if row else [None] * 12
    values = [str(v) if v is not None else None for v in payload.months]
    if row is None:
        row = AssortmentForecast(plan_id=plan_id, base_id=base_id, mt_code=mt)
        db.add(row)
    row.months = json.dumps(values)
    plan.version += 1
    audit(
        db,
        request,
        "forecast",
        {"plan_id": plan_id, "base_id": base_id, "mt": mt, "before": before, "after": values},
    )
    commit(db)
    return {"version": plan.version}


class DisplayInput(BaseModel):
    page_size: Literal[0, 25, 50, 100]


@router.put("/display")
def display(payload: DisplayInput, request: Request, db: Db):
    set_setting(
        db,
        "assortment_page_size",
        str(payload.page_size),
        secret=False,
        actor=actor_from_session(db, "development-admin"),
    )
    audit(db, request, "display", payload.model_dump())
    commit(db)
    return {"pageSize": payload.page_size}


@router.get("/report")
def report(
    db: Db,
    year: int = Query(default=2027, ge=2002, le=2100),
    plan_id: str | None = None,
    mt: ActiveReportingModernTradeCode | None = None,
    search: str = Query(default="", max_length=200),
    mode: Literal["sales", "inventory"] = "sales",
    basis: Literal["net", "gross"] = "net",
    page: int = Query(default=1, ge=1),
    group: str = "",
    model: str = "",
    download: bool = False,
    details: bool = True,
):
    if plan_id:
        plan = db.get(AssortmentPlan, plan_id)
        if not plan or plan.year != year:
            raise HTTPException(422, "แผนและปีไม่ตรงกัน")
    result = assortment.report(
        db, year, plan_id, mt, search, mode, basis, page, download, group, model
    )
    if not download:
        return result
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Assortment"
    identity = ["Base Item", *(assortment.ATTRS if details else []), "Description"]
    sheet.append([f"Assortment {mode} · {basis} · {year}"])
    sheet.append(["แผน: " + (plan.name if plan_id else "—") + " · ยอดเท่าที่มี; ดู Coverage"])
    sheet.append(
        identity
        + [
            f"{code} {label}"
            for code in result["mts"]
            for label in (
                [str(year - 2), str(year - 1), "วันที่ Stock"]
                if mode == "inventory"
                else [
                    str(year - 2),
                    str(year - 1),
                    f"{year} Forecast",
                    f"Growth {year - 1}",
                    "เดือนที่กรอก",
                ]
            )
        ]
    )
    monthly = workbook.create_sheet("Forecast") if mode == "sales" else None
    if monthly:
        monthly.append(identity + ["MT", "ปี", *[str(i) for i in range(1, 13)], "รวมปี", "เดือนที่กรอก"])
    for base in result["items"]:
        prefix = [base["id"], *(base["attributes"] if details else []), base["description"]]
        row = list(prefix)
        for code in result["mts"]:
            cell = base["cells"][code]
            a, b = (cell["years"].get(str(y), {}) for y in (year - 2, year - 1))
            row += [a.get("total"), b.get("total")]
            if mode == "inventory":
                row += [f"{a.get('snapshot') or '—'} / {b.get('snapshot') or '—'}"]
            else:
                row += [cell["forecastTotal"], cell["growth"], f"{cell['filled']}/12"]
                monthly.append(
                    prefix + [code, year, *cell["forecast"], cell["forecastTotal"], cell["filled"]]
                )
        sheet.append(row)
    cov = workbook.create_sheet("Coverage")
    cov.append(["MT", "ปี", "วันที่มีข้อมูล", "วันทั้งปี"])
    for code, years in result["coverage"].items():
        for y, value in years.items():
            cov.append([code, int(y), value["days"], value["expected"]])
    for ws in workbook:
        header = 3 if ws == sheet else 1
        ws.freeze_panes = (
            f"{get_column_letter(len(identity) + 1) if ws != cov else 'B'}{header + 1}"
        )
        ws.auto_filter.ref = f"A{header}:{get_column_letter(ws.max_column)}{ws.max_row}"
        for cell in ws[header]:
            cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="102A43")
        for row in ws:
            for cell in row:
                if cell.data_type == "f":  # User-controlled descriptions must remain text.
                    cell.data_type = "s"
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 45 if i == len(identity) else 18
    output = BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Assortment_{year}_{mode}.xlsx"'},
    )
