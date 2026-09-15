from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

STORAGE_SCALE = Decimal("0.000000000001")
QTY_SCALE = Decimal("0.0001")
SALE_FILENAME = re.compile(
    r"รายงานยอดขาย(?P<from>\d{2}-\d{2}-\d{4})_(?P<to>\d{2}-\d{2}-\d{4})\.xlsx",
    re.IGNORECASE,
)
STOCK_FILENAME = re.compile(
    r"รายงานสต็อคAll(?P<from>\d{2}-\d{2}-\d{4})_(?P<to>\d{2}-\d{2}-\d{4})\.xlsx",
    re.IGNORECASE,
)


class DhFormatError(ValueError):
    """Raised when a DoHome workbook pair violates the source contract."""


@dataclass(frozen=True)
class DhFactRow:
    branch_code: str
    branch_name: str
    sku: str
    description: str | None
    sales_qty: Decimal = Decimal("0")
    stock_on_hand: Decimal = Decimal("0")


@dataclass(frozen=True)
class DhBranchAmount:
    branch_code: str
    branch_name: str
    amount: Decimal


@dataclass(frozen=True)
class DhSummary:
    row_count: int
    store_count: int
    sku_count: int
    negative_row_count: int
    source_amount: Decimal
    amount: Decimal
    sales_qty: Decimal
    stock_on_hand: Decimal


@dataclass(frozen=True)
class DhPairExtract:
    batch_date: date
    sales_date: date
    stock_date: date
    inventory_path: str
    sales_path: str
    inventory_filename: str
    sales_filename: str
    inventory_checksum: str
    sales_checksum: str
    business_fingerprint: str
    sales_rows: tuple[DhFactRow, ...]
    stock_rows: tuple[DhFactRow, ...]
    branch_amounts: tuple[DhBranchAmount, ...]
    sales_summary: DhSummary
    stock_summary: DhSummary
    source_skus: frozenset[str]
    source_branches: tuple[tuple[str, str], ...]
    reconciliation_errors: tuple[str, ...]


def inspect_dh_workbook(source_path: str | Path) -> tuple[str, date]:
    path = Path(source_path)
    kind, data_date = _filename_details(path.name)
    if path.suffix.lower() != ".xlsx":
        raise DhFormatError("DoHome รองรับไฟล์ .xlsx เท่านั้น")
    return kind, data_date


def extract_dh_pair(inventory_path: str | Path, sales_path: str | Path) -> DhPairExtract:
    inventory = Path(inventory_path)
    sales = Path(sales_path)
    inventory_kind, stock_date = inspect_dh_workbook(inventory)
    sales_kind, sales_date = inspect_dh_workbook(sales)
    if inventory_kind != "inventory" or sales_kind != "sales":
        raise DhFormatError("กรุณาเลือกไฟล์ Stock และ Sale ของ DoHome ให้ถูกประเภท")
    if sales_date != stock_date - timedelta(days=1):
        raise DhFormatError("วันที่ Sale ต้องอยู่ก่อน Stock 1 วันใน Batch เดียวกัน")

    stock_rows, branch_code_by_name = _read_stock(inventory)
    sales_rows, branch_amounts, sale_branches, errors = _read_sale(
        sales, branch_code_by_name
    )
    all_skus = {row.sku for row in stock_rows} | {row.sku for row in sales_rows}
    all_branches = {
        (row.branch_code, row.branch_name) for row in (*stock_rows, *sales_rows)
    } | set(sale_branches)

    sales_qty = sum((row.sales_qty for row in sales_rows), Decimal("0"))
    source_amount = sum((row.amount for row in branch_amounts), Decimal("0"))
    stock_on_hand = sum((row.stock_on_hand for row in stock_rows), Decimal("0"))
    sales_summary = DhSummary(
        row_count=len(sales_rows),
        store_count=len(sale_branches),
        sku_count=len({row.sku for row in sales_rows}),
        negative_row_count=sum(row.sales_qty < 0 for row in sales_rows),
        source_amount=source_amount,
        amount=source_amount.quantize(STORAGE_SCALE),
        sales_qty=sales_qty,
        stock_on_hand=Decimal("0"),
    )
    stock_summary = DhSummary(
        row_count=len(stock_rows),
        store_count=len({row.branch_code for row in stock_rows}),
        sku_count=len({row.sku for row in stock_rows}),
        negative_row_count=sum(row.stock_on_hand < 0 for row in stock_rows),
        source_amount=Decimal("0"),
        amount=Decimal("0"),
        sales_qty=Decimal("0"),
        stock_on_hand=stock_on_hand,
    )
    fingerprint_payload = {
        "batch_date": stock_date.isoformat(),
        "sales_date": sales_date.isoformat(),
        "stock_date": stock_date.isoformat(),
        "sales": [
            [row.branch_code, row.sku, str(row.sales_qty)] for row in sales_rows
        ],
        "amounts": [
            [row.branch_code, str(row.amount)] for row in branch_amounts
        ],
        "stock": [
            [row.branch_code, row.sku, str(row.stock_on_hand)] for row in stock_rows
        ],
    }
    fingerprint = hashlib.sha256(
        json.dumps(fingerprint_payload, ensure_ascii=False, separators=(",", ":")).encode(
            "utf-8"
        )
    ).hexdigest()
    return DhPairExtract(
        batch_date=stock_date,
        sales_date=sales_date,
        stock_date=stock_date,
        inventory_path=str(inventory),
        sales_path=str(sales),
        inventory_filename=inventory.name,
        sales_filename=sales.name,
        inventory_checksum=_sha256(inventory),
        sales_checksum=_sha256(sales),
        business_fingerprint=fingerprint,
        sales_rows=tuple(sales_rows),
        stock_rows=tuple(stock_rows),
        branch_amounts=tuple(branch_amounts),
        sales_summary=sales_summary,
        stock_summary=stock_summary,
        source_skus=frozenset(all_skus),
        source_branches=tuple(sorted(all_branches)),
        reconciliation_errors=tuple(errors),
    )


def _read_stock(path: Path) -> tuple[list[DhFactRow], dict[str, str]]:
    _validate_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_column < 8 or sheet.max_row < 2:
            raise DhFormatError(f"โครงสร้างไฟล์ Stock ไม่ครบ 8 columns: {path.name}")
        rows: list[DhFactRow] = []
        branch_code_by_name: dict[str, str] = {}
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            branch_code = _text(values[0])
            branch_name = _text(values[1])
            sku = _sku(values[2])
            if not branch_code or not branch_name or not sku:
                continue
            existing = branch_code_by_name.setdefault(branch_name, branch_code)
            if existing != branch_code:
                raise DhFormatError(
                    f"ชื่อสาขา {branch_name!r} มีมากกว่าหนึ่งรหัสในไฟล์ Stock"
                )
            rows.append(
                DhFactRow(
                    branch_code=branch_code,
                    branch_name=branch_name,
                    sku=sku,
                    description=_text(values[3]) or None,
                    stock_on_hand=_decimal(values[6], row_number, 7),
                )
            )
        if not rows:
            raise DhFormatError(f"ไม่พบข้อมูล Stock ในไฟล์ {path.name}")
        return rows, branch_code_by_name
    finally:
        workbook.close()


def _read_sale(
    path: Path, branch_code_by_name: dict[str, str]
) -> tuple[
    list[DhFactRow],
    list[DhBranchAmount],
    tuple[tuple[str, str], ...],
    list[str],
]:
    _validate_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_column < 5 or sheet.max_row < 3:
            raise DhFormatError(f"โครงสร้างไฟล์ Sale ไม่ครบ: {path.name}")
        values = list(sheet.iter_rows(values_only=True))
        headers = values[0]
        detail_rows = values[1:-1]
        footer = values[-1]
        branch_columns: list[tuple[int, str, str]] = []
        for index in range(3, len(headers) - 1):
            branch_name = _text(headers[index])
            if not branch_name:
                continue
            branch_code = branch_code_by_name.get(branch_name) or _derived_branch_code(
                branch_name
            )
            branch_columns.append((index, branch_code, branch_name))
        if not branch_columns:
            raise DhFormatError(f"ไม่พบ Branch columns ในไฟล์ {path.name}")

        rows: list[DhFactRow] = []
        errors: list[str] = []
        detail_total = Decimal("0")
        for row_number, row in enumerate(detail_rows, start=2):
            sku = _sku(row[0])
            if not sku:
                continue
            row_total = Decimal("0")
            for index, branch_code, branch_name in branch_columns:
                qty = _decimal(row[index], row_number, index + 1)
                row_total += qty
                if qty:
                    rows.append(
                        DhFactRow(
                            branch_code=branch_code,
                            branch_name=branch_name,
                            sku=sku,
                            description=_text(row[1]) or None,
                            sales_qty=qty,
                        )
                    )
            reported_row_total = _decimal(row[-1], row_number, len(row))
            if (row_total - reported_row_total).copy_abs() > QTY_SCALE:
                errors.append(
                    f"QTY แถว {row_number} รวม {row_total} ไม่ตรงกับ Total {reported_row_total}"
                )
            detail_total += row_total

        reported_total = _decimal(footer[-1], len(values), len(footer))
        if (detail_total - reported_total).copy_abs() > QTY_SCALE:
            errors.append(
                f"QTY รวมจากรายการ {detail_total} ไม่ตรงกับ Total {reported_total}"
            )
        amounts = [
            DhBranchAmount(
                branch_code=branch_code,
                branch_name=branch_name,
                amount=_decimal(footer[index], len(values), index + 1).quantize(
                    STORAGE_SCALE
                ),
            )
            for index, branch_code, branch_name in branch_columns
        ]
        return (
            rows,
            amounts,
            tuple((code, name) for _, code, name in branch_columns),
            errors,
        )
    finally:
        workbook.close()


def _filename_details(filename: str) -> tuple[str, date]:
    for kind, pattern in (("sales", SALE_FILENAME), ("inventory", STOCK_FILENAME)):
        match = pattern.fullmatch(filename)
        if not match:
            continue
        from_date = _parse_date(match.group("from"), filename)
        to_date = _parse_date(match.group("to"), filename)
        if from_date != to_date:
            raise DhFormatError(f"ช่วงวันที่ในชื่อไฟล์ DH ต้องเป็นวันเดียวกัน: {filename}")
        return kind, from_date
    raise DhFormatError(
        "ชื่อไฟล์ DH ต้องเป็นรายงานยอดขายDD-MM-YYYY_DD-MM-YYYY.xlsx "
        "หรือรายงานสต็อคAllDD-MM-YYYY_DD-MM-YYYY.xlsx"
    )


def _parse_date(value: str, filename: str) -> date:
    day, month, year = value.split("-")
    try:
        return date(int(year), int(month), int(day))
    except ValueError as exc:
        raise DhFormatError(f"วันที่ในชื่อไฟล์ DH ไม่ถูกต้อง: {filename}") from exc


def _derived_branch_code(branch_name: str) -> str:
    digest = hashlib.sha1(branch_name.encode("utf-8"), usedforsecurity=False).hexdigest()[:12]
    return f"DHN-{digest}"


def _validate_file(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise DhFormatError("DoHome รองรับไฟล์ .xlsx เท่านั้น")
    if not path.is_file() or path.stat().st_size == 0:
        raise DhFormatError(f"ไม่พบไฟล์หรือไฟล์ว่าง: {path.name}")


def _sku(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _text(value: object) -> str:
    return str(value or "").strip()


def _decimal(value: object, row: int, column: int) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation as exc:
        raise DhFormatError(
            f"ค่าตัวเลขไม่ถูกต้องที่แถว {row} column {column}: {value!r}"
        ) from exc


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
