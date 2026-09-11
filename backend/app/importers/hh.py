from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

STORAGE_SCALE = Decimal("0.000000000001")
DATE_PATTERNS = (
    re.compile(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b"),
    re.compile(r"\b(\d{1,2})-(\d{1,2})-(\d{4})\b"),
)
BRANCHES = (
    (7, "HH-UBON", "อุบลราชธานี"),
    (9, "HH-CHAYANGKUN", "ชยางกูร"),
    (11, "HH-WARIN", "วารินฯ"),
    (13, "HH-KHONKAEN", "ขอนแก่น"),
    (15, "HH-AMNAT", "อำนาจ"),
)


class HhFormatError(ValueError):
    """Raised when a HomeHub workbook pair violates the source contract."""


@dataclass(frozen=True)
class HhRow:
    branch_code: str
    branch_name: str
    sku: str
    description: str | None
    category: str | None
    subcategory: str | None
    source_amount: Decimal
    amount: Decimal
    sales_qty: Decimal
    stock_on_hand: Decimal
    stock_value: Decimal


@dataclass(frozen=True)
class HhSummary:
    row_count: int
    store_count: int
    sku_count: int
    negative_row_count: int
    source_amount: Decimal
    amount: Decimal
    sales_qty: Decimal
    stock_on_hand: Decimal
    stock_value: Decimal


@dataclass(frozen=True)
class HhPairExtract:
    data_date: date
    inventory_path: str
    sales_path: str
    inventory_filename: str
    sales_filename: str
    inventory_checksum: str
    sales_checksum: str
    business_fingerprint: str
    rows: tuple[HhRow, ...]
    summary: HhSummary
    inventory_skus: frozenset[str]
    inventory_branches: frozenset[str]
    reconciliation_errors: tuple[str, ...]


@dataclass
class _Values:
    description: str | None = None
    category: str | None = None
    subcategory: str | None = None
    source_amount: Decimal = Decimal("0")
    sales_qty: Decimal = Decimal("0")
    stock_on_hand: Decimal = Decimal("0")
    stock_value: Decimal = Decimal("0")


def inspect_hh_workbook(source_path: str | Path) -> tuple[str, date]:
    path = Path(source_path)
    filename = path.name.lower()
    kind = (
        "inventory"
        if filename.endswith("stockreport.xlsx")
        else "sales"
        if filename.endswith("salereport.xlsx")
        else None
    )
    if kind is None:
        raise HhFormatError("ชื่อไฟล์ HH ต้องระบุ StockReport.xlsx หรือ SaleReport.xlsx")
    data_date, _, _ = _read_matrix(path, kind)
    return kind, data_date


def extract_hh_pair(inventory_path: str | Path, sales_path: str | Path) -> HhPairExtract:
    inventory = Path(inventory_path)
    sales = Path(sales_path)
    inventory_date, inventory_values, inventory_skus = _read_matrix(inventory, "inventory")
    sales_date, sales_values, sales_skus = _read_matrix(sales, "sales")
    if inventory_date != sales_date:
        raise HhFormatError(
            "วันที่ข้อมูล Stock และ Sale ไม่ตรงกัน: "
            f"{inventory_date:%d/%m/%Y} != {sales_date:%d/%m/%Y}"
        )

    all_skus = inventory_skus | sales_skus
    rows: list[HhRow] = []
    for sku in sorted(all_skus):
        for _, branch_code, branch_name in BRANCHES:
            stock = inventory_values.get((branch_code, sku), _Values())
            sale = sales_values.get((branch_code, sku), _Values())
            rows.append(
                HhRow(
                    branch_code=branch_code,
                    branch_name=branch_name,
                    sku=sku,
                    description=sale.description or stock.description,
                    category=sale.category or stock.category,
                    subcategory=sale.subcategory or stock.subcategory,
                    source_amount=sale.source_amount,
                    amount=sale.source_amount.quantize(STORAGE_SCALE),
                    sales_qty=sale.sales_qty,
                    stock_on_hand=stock.stock_on_hand,
                    stock_value=stock.stock_value,
                )
            )

    fingerprint_payload = [
        [
            row.branch_code,
            row.sku,
            str(row.source_amount),
            str(row.sales_qty),
            str(row.stock_on_hand),
            str(row.stock_value),
        ]
        for row in rows
    ]
    fingerprint = hashlib.sha256(
        json.dumps(
            [inventory_date.isoformat(), fingerprint_payload],
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    summary = HhSummary(
        row_count=len(rows),
        store_count=len(BRANCHES),
        sku_count=len(all_skus),
        negative_row_count=sum(
            1 for row in rows if row.source_amount < 0 or row.sales_qty < 0
        ),
        source_amount=sum((row.source_amount for row in rows), Decimal("0")),
        amount=sum((row.amount for row in rows), Decimal("0")),
        sales_qty=sum((row.sales_qty for row in rows), Decimal("0")),
        stock_on_hand=sum((row.stock_on_hand for row in rows), Decimal("0")),
        stock_value=sum((row.stock_value for row in rows), Decimal("0")),
    )
    return HhPairExtract(
        data_date=inventory_date,
        inventory_path=str(inventory),
        sales_path=str(sales),
        inventory_filename=inventory.name,
        sales_filename=sales.name,
        inventory_checksum=_sha256(inventory),
        sales_checksum=_sha256(sales),
        business_fingerprint=fingerprint,
        rows=tuple(rows),
        summary=summary,
        inventory_skus=frozenset(inventory_skus),
        inventory_branches=frozenset(branch[1] for branch in BRANCHES),
        reconciliation_errors=(),
    )


def _read_matrix(
    path: Path, kind: str
) -> tuple[date, dict[tuple[str, str], _Values], set[str]]:
    if path.suffix.lower() != ".xlsx":
        raise HhFormatError("HomeHub รองรับไฟล์ .xlsx เท่านั้น")
    if not path.is_file() or path.stat().st_size == 0:
        raise HhFormatError(f"ไม่พบไฟล์หรือไฟล์ว่าง: {path.name}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_column < 18 or sheet.max_row < 4:
            raise HhFormatError(f"โครงสร้างไฟล์ {path.name} ไม่ครบ 18 columns")
        data_date = _extract_date(str(sheet.cell(1, 1).value or ""))
        values: dict[tuple[str, str], _Values] = defaultdict(_Values)
        skus: set[str] = set()
        for row_number in range(4, sheet.max_row + 1):
            sku = _sku_text(sheet.cell(row_number, 4))
            if not sku:
                continue
            skus.add(sku)
            description = _text(sheet.cell(row_number, 5).value) or None
            category = _text(sheet.cell(row_number, 2).value) or None
            subcategory = _text(sheet.cell(row_number, 3).value) or None
            for qty_column, branch_code, _ in BRANCHES:
                qty = _decimal(sheet.cell(row_number, qty_column).value, row_number, qty_column)
                value = _decimal(
                    sheet.cell(row_number, qty_column + 1).value,
                    row_number,
                    qty_column + 1,
                )
                target = values[(branch_code, sku)]
                target.description = description
                target.category = category
                target.subcategory = subcategory
                if kind == "inventory":
                    target.stock_on_hand += qty
                    target.stock_value += value
                else:
                    target.sales_qty += qty
                    target.source_amount += value
        if not skus:
            raise HhFormatError(f"ไม่พบ SKU ในไฟล์ {path.name}")
        return data_date, values, skus
    finally:
        workbook.close()


def _extract_date(text: str) -> date:
    candidates: list[tuple[int, date]] = []
    for index, pattern in enumerate(DATE_PATTERNS):
        for match in pattern.finditer(text):
            value = (
                date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
                if index == 0
                else date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
            )
            candidates.append((match.start(), value))
    if candidates:
        return min(candidates, key=lambda item: item[0])[1]
    raise HhFormatError("ไม่พบวันที่ข้อมูลใน Header")


def _sku_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, int):
        number_format = str(cell.number_format or "")
        if number_format and set(number_format) == {"0"}:
            return str(value).zfill(len(number_format))
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _text(value: object) -> str:
    return str(value or "").strip()


def _decimal(value: object, row: int, column: int) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation as exc:
        raise HhFormatError(
            f"ค่าตัวเลขไม่ถูกต้องที่แถว {row} column {column}: {value!r}"
        ) from exc


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
