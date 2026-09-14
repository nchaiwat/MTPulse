from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from app.sales_grain import SALES_GRAIN_DAILY, flat_sales_grain

STORAGE_SCALE = Decimal("0.000000000001")
VAT_DIVISOR = Decimal("1.07")
RECONCILIATION_TOLERANCE = Decimal("0.02")
FILENAME_PATTERN = re.compile(
    r"^Runglawan-(\d{4})-(\d{2})-(\d{2})(?:[-_].*)?\.xlsx$",
    re.IGNORECASE,
)
LEGACY_BRANCH_PATTERN = re.compile(r"^GH-\d+$", re.IGNORECASE)
HEADERS = (
    "Product Code",
    "Product Name",
    "Product Status",
    "Branch Code",
    "Branch Name",
    "Stock (Qty)",
    "Stock Amount (Ex VAT)",
    "Stock Amount (In VAT)",
    "Cost (Ex VAT)",
    "Cost (In VAT)",
    "Sale Quantity",
    "Sale Amount",
)


class TaFormatError(ValueError):
    """Raised when a Thai-Aust workbook violates the confirmed source contract."""


@dataclass(frozen=True)
class TaRow:
    branch_code: str
    branch_name: str
    sku: str
    description: str | None
    product_status: str | None
    source_amount: Decimal
    amount: Decimal
    sales_qty: Decimal
    stock_on_hand: Decimal
    stock_value: Decimal
    stock_amount_in_vat: Decimal
    unit_cost_ex_vat: Decimal
    unit_cost_in_vat: Decimal


@dataclass(frozen=True)
class TaSummary:
    row_count: int
    store_count: int
    sku_count: int
    negative_row_count: int
    source_amount: Decimal
    amount: Decimal
    sales_qty: Decimal
    stock_on_hand: Decimal
    stock_value: Decimal


@dataclass(frozen=True)
class TaExtract:
    data_date: date
    sales_grain: str
    sales_window_days: int | None
    source_path: str
    source_filename: str
    checksum_sha256: str
    business_fingerprint: str
    rows: tuple[TaRow, ...]
    summary: TaSummary
    inventory_skus: frozenset[str]
    inventory_branches: frozenset[str]
    product_status_counts: tuple[tuple[str, int], ...]
    footer_totals: tuple[tuple[str, Decimal], ...]
    reconciliation_errors: tuple[str, ...]


def inspect_ta_workbook(source_path: str | Path) -> tuple[str, date]:
    extract = extract_ta_file(source_path)
    return "combined", extract.data_date


def extract_ta_file(source_path: str | Path) -> TaExtract:
    path = Path(source_path)
    data_date = _data_date_from_filename(path.name)
    if not path.is_file() or path.stat().st_size == 0:
        raise TaFormatError(f"ไม่พบไฟล์หรือไฟล์ว่าง: {path.name}")

    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook.active
        actual_headers = tuple(_text(cell.value) for cell in sheet[1][: len(HEADERS)])
        if actual_headers != HEADERS:
            return _extract_ta_legacy_file(path, data_date, sheet)
        sales_grain, sales_window_days = flat_sales_grain(data_date)

        rows: list[TaRow] = []
        keys: set[tuple[str, str]] = set()
        statuses: Counter[str] = Counter()
        footer: dict[str, Decimal] | None = None
        warnings: list[str] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=12), start=2):
            values = [cell.value for cell in cells]
            sku = _identifier(cells[0])
            branch_code = _identifier(cells[3])
            if not sku and not branch_code:
                if not any(value not in (None, "") for value in values):
                    continue
                if footer is not None:
                    raise TaFormatError("พบ Footer มากกว่าหนึ่งแถว")
                footer = {
                    "Stock Amount (Ex VAT)": _decimal(values[6], row_number, 7),
                    "Stock Amount (In VAT)": _decimal(values[7], row_number, 8),
                    "Cost (Ex VAT)": _decimal(values[8], row_number, 9),
                    "Cost (In VAT)": _decimal(values[9], row_number, 10),
                    "Sale Amount": _decimal(values[11], row_number, 12),
                }
                continue
            if not sku or not branch_code:
                raise TaFormatError(f"ข้อมูล TA แถว {row_number} ขาด Product Code หรือ Branch Code")
            key = (branch_code, sku)
            if key in keys:
                raise TaFormatError(f"พบ SKU × Branch ซ้ำที่แถว {row_number}: {sku} × {branch_code}")
            keys.add(key)

            description = _text(values[1]) or None
            product_status = _text(values[2]) or None
            branch_name = _text(values[4])
            if not description or not branch_name:
                raise TaFormatError(f"ข้อมูล TA แถว {row_number} ขาด Product Name หรือ Branch Name")
            if product_status:
                statuses[product_status] += 1

            stock_qty = _decimal(values[5], row_number, 6)
            stock_ex = _decimal(values[6], row_number, 7)
            stock_in = _decimal(values[7], row_number, 8)
            cost_ex = _decimal(values[8], row_number, 9)
            cost_in = _decimal(values[9], row_number, 10)
            sales_qty = _decimal(values[10], row_number, 11)
            source_amount = _decimal(values[11], row_number, 12)
            _compare(
                warnings,
                f"แถว {row_number} Stock Amount (Ex VAT)",
                stock_ex,
                stock_qty * cost_ex,
            )
            rows.append(
                TaRow(
                    branch_code=branch_code,
                    branch_name=branch_name,
                    sku=sku,
                    description=description,
                    product_status=product_status,
                    source_amount=source_amount,
                    amount=(source_amount / VAT_DIVISOR).quantize(STORAGE_SCALE),
                    sales_qty=sales_qty,
                    stock_on_hand=stock_qty,
                    stock_value=stock_ex,
                    stock_amount_in_vat=stock_in,
                    unit_cost_ex_vat=cost_ex,
                    unit_cost_in_vat=cost_in,
                )
            )
    finally:
        workbook.close()

    if not rows:
        raise TaFormatError("ไม่พบข้อมูล Detail ในไฟล์ TA")
    if footer is None:
        raise TaFormatError("ไม่พบ Footer สำหรับตรวจสอบยอดรวมของ TA")

    detail_totals = {
        "Stock Amount (Ex VAT)": sum((row.stock_value for row in rows), Decimal("0")),
        "Stock Amount (In VAT)": sum((row.stock_amount_in_vat for row in rows), Decimal("0")),
        "Cost (Ex VAT)": sum((row.unit_cost_ex_vat for row in rows), Decimal("0")),
        "Cost (In VAT)": sum((row.unit_cost_in_vat for row in rows), Decimal("0")),
        "Sale Amount": sum((row.source_amount for row in rows), Decimal("0")),
    }
    for label, expected in footer.items():
        _compare(warnings, f"Footer {label}", detail_totals[label], expected)

    fingerprint_rows = [
        [
            row.branch_code,
            row.branch_name,
            row.sku,
            row.description,
            row.product_status,
            str(row.source_amount),
            str(row.sales_qty),
            str(row.stock_on_hand),
            str(row.stock_value),
            str(row.stock_amount_in_vat),
            str(row.unit_cost_ex_vat),
            str(row.unit_cost_in_vat),
        ]
        for row in sorted(rows, key=lambda item: (item.branch_code, item.sku))
    ]
    fingerprint = hashlib.sha256(
        json.dumps(
            [data_date.isoformat(), fingerprint_rows],
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    summary = TaSummary(
        row_count=len(rows),
        store_count=len({row.branch_code for row in rows}),
        sku_count=len({row.sku for row in rows}),
        negative_row_count=sum(row.source_amount < 0 or row.sales_qty < 0 for row in rows),
        source_amount=detail_totals["Sale Amount"],
        amount=sum((row.amount for row in rows), Decimal("0")),
        sales_qty=sum((row.sales_qty for row in rows), Decimal("0")),
        stock_on_hand=sum((row.stock_on_hand for row in rows), Decimal("0")),
        stock_value=detail_totals["Stock Amount (Ex VAT)"],
    )
    return TaExtract(
        data_date=data_date,
        sales_grain=sales_grain,
        sales_window_days=sales_window_days,
        source_path=str(path),
        source_filename=path.name,
        checksum_sha256=_sha256(path),
        business_fingerprint=fingerprint,
        rows=tuple(rows),
        summary=summary,
        inventory_skus=frozenset(row.sku for row in rows),
        inventory_branches=frozenset(row.branch_code for row in rows),
        product_status_counts=tuple(sorted(statuses.items())),
        footer_totals=tuple(footer.items()),
        reconciliation_errors=tuple(warnings),
    )


def _extract_ta_legacy_file(path: Path, data_date: date, sheet) -> TaExtract:
    allow_empty = _is_empty_legacy_sheet(sheet)
    rows, statuses, footer, warnings = _extract_legacy_rows(sheet)
    if not rows and not allow_empty:
        raise TaFormatError("ไม่พบข้อมูล Detail ในไฟล์ TA")

    detail_totals = {
        "Stock (Qty)": sum((row.stock_on_hand for row in rows), Decimal("0")),
        "Sale Quantity": sum((row.sales_qty for row in rows), Decimal("0")),
        "Sale Amount": sum((row.source_amount for row in rows), Decimal("0")),
    }
    for label, expected in footer.items():
        _compare(warnings, f"Footer {label}", detail_totals[label], expected)

    fingerprint_rows = [
        [
            row.branch_code,
            row.branch_name,
            row.sku,
            row.description,
            row.product_status,
            str(row.source_amount),
            str(row.sales_qty),
            str(row.stock_on_hand),
        ]
        for row in sorted(rows, key=lambda item: (item.branch_code, item.sku))
    ]
    fingerprint = hashlib.sha256(
        json.dumps(
            [data_date.isoformat(), fingerprint_rows],
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    summary = TaSummary(
        row_count=len(rows),
        store_count=len({row.branch_code for row in rows}),
        sku_count=len({row.sku for row in rows}),
        negative_row_count=sum(
            row.source_amount < 0 or row.sales_qty < 0 for row in rows
        ),
        source_amount=detail_totals["Sale Amount"],
        amount=sum((row.amount for row in rows), Decimal("0")),
        sales_qty=detail_totals["Sale Quantity"],
        stock_on_hand=detail_totals["Stock (Qty)"],
        stock_value=Decimal("0"),
    )
    return TaExtract(
        data_date=data_date,
        sales_grain=SALES_GRAIN_DAILY,
        sales_window_days=None,
        source_path=str(path),
        source_filename=path.name,
        checksum_sha256=_sha256(path),
        business_fingerprint=fingerprint,
        rows=tuple(rows),
        summary=summary,
        inventory_skus=frozenset(row.sku for row in rows),
        inventory_branches=frozenset(row.branch_code for row in rows),
        product_status_counts=tuple(sorted(statuses.items())),
        footer_totals=tuple(footer.items()),
        reconciliation_errors=tuple(warnings),
    )


def _extract_legacy_rows(
    sheet,
) -> tuple[list[TaRow], Counter[str], dict[str, Decimal], list[str]]:
    if _is_empty_legacy_sheet(sheet):
        return (
            [],
            Counter(),
            {
                "Stock (Qty)": Decimal("0"),
                "Sale Quantity": Decimal("0"),
                "Sale Amount": Decimal("0"),
            },
            [],
        )
    if sheet.max_column < 12 or _text(sheet.cell(4, 7).value).casefold() != "inactive":
        raise TaFormatError(
            "โครงสร้าง column ของ TA ไม่ถูกต้อง: "
            f"ต้องเป็น {', '.join(HEADERS)} หรือ TA Legacy Branch Matrix"
        )

    branch_starts: list[tuple[int, str, str]] = []
    grand_total_start: int | None = None
    for column in range(8, sheet.max_column + 1):
        value = _text(sheet.cell(2, column).value)
        if LEGACY_BRANCH_PATTERN.fullmatch(value):
            branch_name = _text(sheet.cell(3, column).value)
            if not branch_name:
                raise TaFormatError(f"ข้อมูล TA Legacy column {column} ขาด Branch Name")
            branch_starts.append((column, value, branch_name))
        elif value.casefold() == "grand total":
            grand_total_start = column

    if not branch_starts or grand_total_start is None:
        raise TaFormatError("โครงสร้าง TA Legacy ไม่พบ Branch หรือ Grand Total")

    group_starts = [column for column, _, _ in branch_starts] + [grand_total_start]
    branch_groups: list[tuple[str, str, tuple[int, int, int]]] = []
    for index, (start, branch_code, branch_name) in enumerate(branch_starts):
        end = group_starts[index + 1] - 1
        metric_columns = tuple(
            column
            for column in range(start, end + 1)
            if _text(sheet.cell(4, column).value)
        )
        if len(metric_columns) != 3:
            raise TaFormatError(
                f"โครงสร้าง TA Legacy ของ Branch {branch_code} ต้องมี 3 Metrics"
            )
        branch_groups.append((branch_code, branch_name, metric_columns))

    total_columns = tuple(
        column
        for column in range(grand_total_start, sheet.max_column + 1)
        if _text(sheet.cell(4, column).value)
    )
    if len(total_columns) != 3:
        raise TaFormatError("โครงสร้าง TA Legacy ของ Grand Total ต้องมี 3 Metrics")

    footer_row = next(
        (
            row_number
            for row_number in range(5, sheet.max_row + 1)
            if _text(sheet.cell(row_number, 1).value).casefold() == "grand total"
        ),
        None,
    )
    if footer_row is None:
        raise TaFormatError("ไม่พบ Footer สำหรับตรวจสอบยอดรวมของ TA Legacy")

    rows: list[TaRow] = []
    keys: set[tuple[str, str]] = set()
    statuses: Counter[str] = Counter()
    warnings: list[str] = []
    branch_totals = {
        branch_code: [Decimal("0"), Decimal("0"), Decimal("0")]
        for branch_code, _, _ in branch_groups
    }
    for row_number in range(5, footer_row):
        sku = _identifier(sheet.cell(row_number, 1))
        if not sku:
            continue
        description = _text(sheet.cell(row_number, 3).value) or None
        product_status = _text(sheet.cell(row_number, 7).value) or None
        if not description:
            raise TaFormatError(f"ข้อมูล TA Legacy แถว {row_number} ขาด Product Name")
        for branch_code, branch_name, metric_columns in branch_groups:
            stock_qty, sales_qty, source_amount = (
                _decimal(sheet.cell(row_number, column).value, row_number, column)
                for column in metric_columns
            )
            if stock_qty == sales_qty == source_amount == 0:
                continue
            key = (branch_code, sku)
            if key in keys:
                raise TaFormatError(
                    f"พบ SKU × Branch ซ้ำที่แถว {row_number}: {sku} × {branch_code}"
                )
            keys.add(key)
            if product_status:
                statuses[product_status] += 1
            totals = branch_totals[branch_code]
            totals[0] += stock_qty
            totals[1] += sales_qty
            totals[2] += source_amount
            rows.append(
                TaRow(
                    branch_code=branch_code,
                    branch_name=branch_name,
                    sku=sku,
                    description=description,
                    product_status=product_status,
                    source_amount=source_amount,
                    amount=(source_amount / VAT_DIVISOR).quantize(STORAGE_SCALE),
                    sales_qty=sales_qty,
                    stock_on_hand=stock_qty,
                    stock_value=Decimal("0"),
                    stock_amount_in_vat=Decimal("0"),
                    unit_cost_ex_vat=Decimal("0"),
                    unit_cost_in_vat=Decimal("0"),
                )
            )

    for branch_code, _, metric_columns in branch_groups:
        for label, column, actual in zip(
            ("Stock (Qty)", "Sale Quantity", "Sale Amount"),
            metric_columns,
            branch_totals[branch_code],
            strict=True,
        ):
            _compare(
                warnings,
                f"Footer {branch_code} {label}",
                actual,
                _decimal(sheet.cell(footer_row, column).value, footer_row, column),
            )

    footer = {
        "Stock (Qty)": _decimal(
            sheet.cell(footer_row, total_columns[0]).value,
            footer_row,
            total_columns[0],
        ),
        "Sale Quantity": _decimal(
            sheet.cell(footer_row, total_columns[1]).value,
            footer_row,
            total_columns[1],
        ),
        "Sale Amount": _decimal(
            sheet.cell(footer_row, total_columns[2]).value,
            footer_row,
            total_columns[2],
        ),
    }
    return rows, statuses, footer, warnings


def _is_empty_legacy_sheet(sheet) -> bool:
    return (
        sheet.max_row == 4
        and _text(sheet.cell(2, 8).value).casefold() == "grand total"
        and _text(sheet.cell(3, 7).value).casefold() == "inactive"
        and _text(sheet.cell(4, 1).value).casefold() == "grand total"
    )


def _data_date_from_filename(filename: str) -> date:
    match = FILENAME_PATTERN.fullmatch(filename)
    if match is None:
        raise TaFormatError("ชื่อไฟล์ TA ต้องขึ้นต้น Runglawan และมีวันที่รูปแบบ Runglawan-YYYY-MM-DD*.xlsx")
    try:
        received_date = date(*(int(part) for part in match.groups()))
    except ValueError as exc:
        raise TaFormatError(f"วันที่ในชื่อไฟล์ TA ไม่ถูกต้อง: {filename}") from exc
    return received_date - timedelta(days=1)


def _identifier(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, int):
        number_format = str(cell.number_format or "")
        if number_format and set(number_format) == {"0"}:
            return str(value).zfill(len(number_format))
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _text(value: object) -> str:
    return str(value or "").strip()


def _decimal(value: object, row: int, column: int) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation as exc:
        raise TaFormatError(f"ค่าตัวเลขไม่ถูกต้องที่แถว {row} column {column}: {value!r}") from exc


def _compare(warnings: list[str], label: str, actual: Decimal, expected: Decimal) -> None:
    if abs(actual - expected) > RECONCILIATION_TOLERANCE:
        warnings.append(f"{label}: ผลรวม Detail {actual} ไม่ตรงกับ Source {expected}")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
