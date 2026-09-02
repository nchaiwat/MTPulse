from __future__ import annotations

import hashlib
import shutil
import tempfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.biffh import XLRDError

SHEET_NAME = "ReportSaleSubscription"
HEADER_ROW = 5
DATA_START_ROW = 6
VAT_DIVISOR = Decimal("1.07")
STORAGE_SCALE = Decimal("0.000000000001")
EXPECTED_HEADERS = {
    0: "Store",
    2: "Cat",
    3: "Sub Cat",
    4: "Brand",
    5: "SKU",
    6: "Barcode",
    7: "Description",
    9: "Product Type",
    10: "Sales Amount",
    11: "Sales Qty",
    12: "Stock On Hand",
    13: "Stock On Order",
    14: "Last Sold Date",
    15: "Last Receive Date",
}


class TwdFormatError(ValueError):
    """Raised when a TWD workbook violates the agreed source contract."""


@dataclass(frozen=True)
class TwdRow:
    branch_code: str
    branch_name: str
    category: str | None
    subcategory: str | None
    brand: str | None
    sku: str
    barcode: str | None
    description: str | None
    product_type: str | None
    source_amount: Decimal
    amount: Decimal
    sales_qty: Decimal
    stock_on_hand: Decimal
    stock_on_order: Decimal
    last_sold_date: date | None
    last_receive_date: date | None


@dataclass(frozen=True)
class TwdSummary:
    row_count: int
    store_count: int
    sku_count: int
    negative_row_count: int
    source_amount: Decimal
    amount: Decimal
    sales_qty: Decimal
    stock_on_hand: Decimal
    stock_on_order: Decimal


@dataclass(frozen=True)
class TwdExtract:
    source_path: str
    source_filename: str
    checksum_sha256: str
    data_date: date
    rows: tuple[TwdRow, ...]
    summary: TwdSummary
    reported_summary: TwdSummary
    reconciliation_errors: tuple[str, ...]


def calculate_amount(source_amount: Decimal) -> Decimal:
    return (source_amount / VAT_DIVISOR).quantize(STORAGE_SCALE)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_twd_file(source_path: str | Path) -> TwdExtract:
    """Copy the immutable source to temp, parse it, then always remove the copy."""
    source = Path(source_path)
    if not source.is_file():
        raise FileNotFoundError(source)
    with tempfile.TemporaryDirectory(prefix="mtpulse-twd-") as temp_dir:
        copied = Path(temp_dir) / source.name
        shutil.copy2(source, copied)
        return _read_copied_file(copied, str(source))


def _read_copied_file(path: Path, original_path: str) -> TwdExtract:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xls":
            return _read_xls_copied_file(path, original_path)
        if suffix == ".xlsx":
            return _read_xlsx_copied_file(path, original_path)
        raise TwdFormatError("รองรับเฉพาะไฟล์ TWD นามสกุล .xls และ .xlsx")
    except TwdFormatError:
        raise
    except (XLRDError, BadZipFile, InvalidFileException, EOFError) as exc:
        raise TwdFormatError(
            f"เปิดไฟล์ {path.name} ไม่สำเร็จ: {exc}"
        ) from exc


def _read_xls_copied_file(path: Path, original_path: str) -> TwdExtract:
    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        if SHEET_NAME not in workbook.sheet_names():
            raise TwdFormatError(f"ไม่พบ sheet '{SHEET_NAME}'")
        sheet = workbook.sheet_by_name(SHEET_NAME)
        _validate_headers(sheet)
        data_date = _parse_period(sheet.cell_value(2, 1))
        source_cell_errors: Counter[tuple[str, str]] = Counter()
        rows = tuple(_read_rows(sheet, workbook.datemode, source_cell_errors))
        reported_summary = _read_reported_summary(sheet)
    finally:
        workbook.release_resources()
    if not rows:
        raise TwdFormatError("ไม่พบรายการข้อมูล")
    _validate_unique_branch_sku(rows)
    summary = summarize(rows)
    return TwdExtract(
        source_path=original_path,
        source_filename=path.name,
        checksum_sha256=sha256_file(path),
        data_date=data_date,
        rows=rows,
        summary=summary,
        reported_summary=reported_summary,
        reconciliation_errors=(
            _source_cell_warnings(source_cell_errors)
            + _reconcile(summary, reported_summary)
        ),
    )


def _read_xlsx_copied_file(path: Path, original_path: str) -> TwdExtract:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise TwdFormatError(f"ไม่พบ sheet '{SHEET_NAME}'")
        sheet = workbook[SHEET_NAME]
        _validate_xlsx_headers(sheet)
        data_date = _parse_period(sheet.cell(row=3, column=2).value)
        source_cell_errors: Counter[tuple[str, str]] = Counter()
        rows = tuple(_read_xlsx_rows(sheet, source_cell_errors))
        reported_summary = _read_xlsx_reported_summary(sheet)
    finally:
        workbook.close()
    if not rows:
        raise TwdFormatError("ไม่พบรายการข้อมูล")
    _validate_unique_branch_sku(rows)
    summary = summarize(rows)
    return TwdExtract(
        source_path=original_path,
        source_filename=path.name,
        checksum_sha256=sha256_file(path),
        data_date=data_date,
        rows=rows,
        summary=summary,
        reported_summary=reported_summary,
        reconciliation_errors=(
            _source_cell_warnings(source_cell_errors)
            + _reconcile(summary, reported_summary)
        ),
    )


def _read_reported_summary(sheet: xlrd.sheet.Sheet) -> TwdSummary:
    total_row = next(
        (
            row
            for row in range(DATA_START_ROW, sheet.nrows)
            if str(sheet.cell_value(row, 7)).strip() == "Total"
        ),
        None,
    )
    if total_row is None:
        raise TwdFormatError("ไม่พบแถว Total")
    count_text = next(
        (
            str(sheet.cell_value(row, 0)).strip()
            for row in range(total_row + 1, sheet.nrows)
            if str(sheet.cell_value(row, 0)).strip().startswith("Count of Rows")
        ),
        "",
    )
    try:
        row_count = int(count_text.rsplit("-", 1)[1].strip())
    except (IndexError, ValueError) as exc:
        raise TwdFormatError(f"Count of Rows ไม่ถูกต้อง: {count_text!r}") from exc
    source_amount = _decimal(sheet.cell_value(total_row, 10))
    return TwdSummary(
        row_count=row_count,
        store_count=0,
        sku_count=0,
        negative_row_count=0,
        source_amount=source_amount,
        amount=calculate_amount(source_amount),
        sales_qty=_decimal(sheet.cell_value(total_row, 11)),
        stock_on_hand=_decimal(sheet.cell_value(total_row, 12)),
        stock_on_order=_decimal(sheet.cell_value(total_row, 13)),
    )


def _reconcile(calculated: TwdSummary, reported: TwdSummary) -> tuple[str, ...]:
    checks = {
        "Row Count": (calculated.row_count, reported.row_count),
        "Sales Amount": (calculated.source_amount, reported.source_amount),
        "Sales Qty": (calculated.sales_qty, reported.sales_qty),
        "Stock On Hand": (calculated.stock_on_hand, reported.stock_on_hand),
        "Stock On Order": (calculated.stock_on_order, reported.stock_on_order),
    }
    return tuple(
        f"{name}: calculated={calculated_value}, source={reported_value}"
        for name, (calculated_value, reported_value) in checks.items()
        if calculated_value != reported_value
    )


def _validate_headers(sheet: xlrd.sheet.Sheet) -> None:
    mismatches = [
        f"column {column + 1}: expected {expected!r}"
        for column, expected in EXPECTED_HEADERS.items()
        if str(sheet.cell_value(HEADER_ROW, column)).strip() != expected
    ]
    if mismatches:
        raise TwdFormatError("โครงสร้าง header ไม่ตรง: " + "; ".join(mismatches))


def _validate_xlsx_headers(sheet: Any) -> None:
    mismatches = [
        f"column {column + 1}: expected {expected!r}"
        for column, expected in EXPECTED_HEADERS.items()
        if str(sheet.cell(row=HEADER_ROW + 1, column=column + 1).value or "").strip()
        != expected
    ]
    if mismatches:
        raise TwdFormatError("โครงสร้าง header ไม่ตรง: " + "; ".join(mismatches))


def _parse_period(value: Any) -> date:
    text = str(value).strip()
    try:
        return datetime.strptime(text, "%a,%d %b %Y").date()
    except ValueError as exc:
        raise TwdFormatError(f"Period ไม่ถูกต้อง: {text!r}") from exc


def _read_rows(
    sheet: xlrd.sheet.Sheet,
    datemode: int,
    source_cell_errors: Counter[tuple[str, str]],
):
    for row_index in range(DATA_START_ROW, sheet.nrows):
        store = str(sheet.cell_value(row_index, 0)).strip()
        if not store or store.startswith("Count of Rows"):
            continue
        branch_code, separator, branch_name = store.partition("-")
        if not separator or not branch_code.strip() or not branch_name.strip():
            raise TwdFormatError(f"Store แถว {row_index + 1} ไม่ถูกต้อง: {store!r}")
        sku = _identifier(sheet.cell_value(row_index, 5))
        if not sku:
            raise TwdFormatError(f"SKU แถว {row_index + 1} ว่าง")
        source_amount = _decimal_cell(
            sheet.cell(row_index, 10), "Sales Amount", source_cell_errors
        )
        yield TwdRow(
            branch_code=branch_code.strip(),
            branch_name=branch_name.strip(),
            category=_identifier(sheet.cell_value(row_index, 2), 3),
            subcategory=_identifier(sheet.cell_value(row_index, 3), 3),
            brand=_text(sheet.cell_value(row_index, 4)),
            sku=sku,
            barcode=_identifier(sheet.cell_value(row_index, 6)),
            description=_text(sheet.cell_value(row_index, 7)),
            product_type=_text(sheet.cell_value(row_index, 9)),
            source_amount=source_amount,
            amount=calculate_amount(source_amount),
            sales_qty=_decimal_cell(
                sheet.cell(row_index, 11), "Sales Qty", source_cell_errors
            ),
            stock_on_hand=_decimal_cell(
                sheet.cell(row_index, 12), "Stock On Hand", source_cell_errors
            ),
            stock_on_order=_decimal_cell(
                sheet.cell(row_index, 13), "Stock On Order", source_cell_errors
            ),
            last_sold_date=_excel_date(sheet.cell(row_index, 14), datemode),
            last_receive_date=_excel_date(sheet.cell(row_index, 15), datemode),
        )


def _read_xlsx_reported_summary(sheet: Any) -> TwdSummary:
    total_row = next(
        (
            row
            for row in range(DATA_START_ROW + 1, sheet.max_row + 1)
            if str(sheet.cell(row=row, column=8).value or "").strip() == "Total"
        ),
        None,
    )
    if total_row is None:
        raise TwdFormatError("ไม่พบแถว Total")
    count_text = next(
        (
            str(sheet.cell(row=row, column=1).value or "").strip()
            for row in range(total_row + 1, sheet.max_row + 1)
            if str(sheet.cell(row=row, column=1).value or "")
            .strip()
            .startswith("Count of Rows")
        ),
        "",
    )
    try:
        row_count = int(count_text.rsplit("-", 1)[1].strip())
    except (IndexError, ValueError) as exc:
        raise TwdFormatError(f"Count of Rows ไม่ถูกต้อง: {count_text!r}") from exc
    source_amount = _decimal(sheet.cell(row=total_row, column=11).value)
    return TwdSummary(
        row_count=row_count,
        store_count=0,
        sku_count=0,
        negative_row_count=0,
        source_amount=source_amount,
        amount=calculate_amount(source_amount),
        sales_qty=_decimal(sheet.cell(row=total_row, column=12).value),
        stock_on_hand=_decimal(sheet.cell(row=total_row, column=13).value),
        stock_on_order=_decimal(sheet.cell(row=total_row, column=14).value),
    )


def _read_xlsx_rows(
    sheet: Any,
    source_cell_errors: Counter[tuple[str, str]],
):
    for row_index in range(DATA_START_ROW + 1, sheet.max_row + 1):
        store = str(sheet.cell(row=row_index, column=1).value or "").strip()
        if not store or store.startswith("Count of Rows"):
            continue
        branch_code, separator, branch_name = store.partition("-")
        if not separator or not branch_code.strip() or not branch_name.strip():
            raise TwdFormatError(f"Store แถว {row_index} ไม่ถูกต้อง: {store!r}")
        sku = _identifier(sheet.cell(row=row_index, column=6).value)
        if not sku:
            raise TwdFormatError(f"SKU แถว {row_index} ว่าง")
        source_amount = _xlsx_decimal_cell(
            sheet.cell(row=row_index, column=11),
            "Sales Amount",
            source_cell_errors,
        )
        yield TwdRow(
            branch_code=branch_code.strip(),
            branch_name=branch_name.strip(),
            category=_identifier(sheet.cell(row=row_index, column=3).value, 3),
            subcategory=_identifier(sheet.cell(row=row_index, column=4).value, 3),
            brand=_text(sheet.cell(row=row_index, column=5).value),
            sku=sku,
            barcode=_identifier(sheet.cell(row=row_index, column=7).value),
            description=_text(sheet.cell(row=row_index, column=8).value),
            product_type=_text(sheet.cell(row=row_index, column=10).value),
            source_amount=source_amount,
            amount=calculate_amount(source_amount),
            sales_qty=_xlsx_decimal_cell(
                sheet.cell(row=row_index, column=12),
                "Sales Qty",
                source_cell_errors,
            ),
            stock_on_hand=_xlsx_decimal_cell(
                sheet.cell(row=row_index, column=13),
                "Stock On Hand",
                source_cell_errors,
            ),
            stock_on_order=_xlsx_decimal_cell(
                sheet.cell(row=row_index, column=14),
                "Stock On Order",
                source_cell_errors,
            ),
            last_sold_date=_xlsx_date(sheet.cell(row=row_index, column=15).value),
            last_receive_date=_xlsx_date(sheet.cell(row=row_index, column=16).value),
        )


def _identifier(value: Any, width: int | None = None) -> str | None:
    if value in (None, ""):
        return None
    result = (
        str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()
    )
    return result.zfill(width) if width else result


def _text(value: Any) -> str | None:
    result = str(value).strip() if value is not None else ""
    return result or None


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise TwdFormatError(f"ค่าตัวเลขไม่ถูกต้อง: {value!r}") from exc


def _decimal_cell(
    cell: xlrd.sheet.Cell,
    field_name: str,
    source_cell_errors: Counter[tuple[str, str]],
) -> Decimal:
    if cell.ctype == xlrd.XL_CELL_ERROR:
        error_text = xlrd.error_text_from_code.get(cell.value, f"error code {cell.value}")
        source_cell_errors[(field_name, error_text)] += 1
        return Decimal("0")
    return _decimal(cell.value)


def _xlsx_decimal_cell(
    cell: Any,
    field_name: str,
    source_cell_errors: Counter[tuple[str, str]],
) -> Decimal:
    if cell.data_type == "e":
        source_cell_errors[(field_name, str(cell.value))] += 1
        return Decimal("0")
    return _decimal(cell.value)


def _source_cell_warnings(
    source_cell_errors: Counter[tuple[str, str]],
) -> tuple[str, ...]:
    return tuple(
        f"{field_name}: พบเซลล์ต้นทาง {error_text} จำนวน {count} เซลล์ ระบบไม่นำมารวมยอด"
        for (field_name, error_text), count in sorted(source_cell_errors.items())
    )


def _excel_date(cell: xlrd.sheet.Cell, datemode: int) -> date | None:
    if cell.value in (None, ""):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, datemode).date()
    try:
        return datetime.strptime(str(cell.value).strip(), "%Y-%m-%d").date()
    except ValueError as exc:
        raise TwdFormatError(f"วันที่ไม่ถูกต้อง: {cell.value!r}") from exc


def _xlsx_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        parsed = from_excel(value)
        return parsed.date() if isinstance(parsed, datetime) else parsed
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except ValueError as exc:
        raise TwdFormatError(f"วันที่ไม่ถูกต้อง: {value!r}") from exc


def _validate_unique_branch_sku(rows: tuple[TwdRow, ...]) -> None:
    seen: set[tuple[str, str]] = set()
    for row in rows:
        key = (row.branch_code, row.sku)
        if key in seen:
            raise TwdFormatError(f"พบ Branch × SKU ซ้ำ: {row.branch_code} × {row.sku}")
        seen.add(key)


def summarize(rows: tuple[TwdRow, ...]) -> TwdSummary:
    return TwdSummary(
        row_count=len(rows),
        store_count=len({row.branch_code for row in rows}),
        sku_count=len({row.sku for row in rows}),
        negative_row_count=sum(row.source_amount < 0 for row in rows),
        source_amount=sum((row.source_amount for row in rows), Decimal("0")),
        amount=sum((row.amount for row in rows), Decimal("0")),
        sales_qty=sum((row.sales_qty for row in rows), Decimal("0")),
        stock_on_hand=sum((row.stock_on_hand for row in rows), Decimal("0")),
        stock_on_order=sum((row.stock_on_order for row in rows), Decimal("0")),
    )
