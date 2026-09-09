from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.local_time import as_bangkok, bangkok_now

PERIOD_LABELS = {
    "ytd": "YTD",
    "h1": "H1",
    "h2": "H2",
    "full": "Full Year",
}
METRIC_LABELS = {"amount": "Amount", "qty": "Qty"}


def dashboard_export_filename(
    mt_code: str,
    year: int,
    period: str,
    metric: str,
    current_time: datetime | None = None,
) -> str:
    timestamp = as_bangkok(current_time or bangkok_now()).strftime("%Y%m%d_%H%M%S")
    return (
        f"{mt_code}_Dashboard_{year}_{PERIOD_LABELS[period]}_"
        f"{METRIC_LABELS[metric]}_{timestamp}.xlsx"
    )


def _percent(value: float | None) -> float | None:
    return None if value is None else value / 100


def build_dashboard_workbook(report: dict, *, metric: str) -> bytes:
    meta = report["meta"]
    summary = report.get("summary")
    year = meta["year"]
    previous_year = meta["previousYear"]
    period = PERIOD_LABELS[meta["period"]]
    mt_code = meta["mtCode"]

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Dashboard"
    sheet.sheet_view.showGridLines = False

    navy = "102A43"
    sky = "02ABFF"
    pale_blue = "EFF5F8"
    line = "DCE6EE"
    muted = "64768A"
    white = "FFFFFF"
    thin_border = Border(bottom=Side(style="thin", color=line))

    sheet.merge_cells("A1:H1")
    sheet["A1"] = f"{mt_code} Sales Dashboard"
    sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    sheet["A1"].font = Font(color=white, bold=True, size=16)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:H2")
    sheet["A2"] = (
        f"Year: {year} | Period: {period} | Metric: {METRIC_LABELS[metric]} | "
        f"Latest data: {meta.get('latestDataDate') or '-'} | "
        f"Export: {bangkok_now():%d/%m/%Y %H:%M:%S}"
    )
    sheet["A2"].font = Font(color=muted, italic=True, size=9)

    kpi_headers = [
        f"Sales Ex.VAT {year}",
        f"Sales Ex.VAT {previous_year}",
        "Sales YoY",
        f"Qty {year}",
        f"Qty {previous_year}",
        "Qty YoY",
        "Data completeness",
        "Loaded / Expected days",
    ]
    kpi_values = [
        summary["currentAmount"] if summary else 0,
        summary["previousAmount"] if summary else 0,
        _percent(summary["amountYoY"]) if summary else None,
        summary["currentQty"] if summary else 0,
        summary["previousQty"] if summary else 0,
        _percent(summary["qtyYoY"]) if summary else None,
        meta["completenessPercent"] / 100,
        f"{meta['loadedDays']} / {meta['expectedDays']}",
    ]
    for column, (header, value) in enumerate(zip(kpi_headers, kpi_values, strict=True), 1):
        header_cell = sheet.cell(4, column, header)
        header_cell.fill = PatternFill("solid", fgColor=pale_blue)
        header_cell.font = Font(color=muted, bold=True, size=9)
        header_cell.alignment = Alignment(vertical="center", wrap_text=True)
        value_cell = sheet.cell(5, column, value)
        value_cell.font = Font(color=navy, bold=True, size=13)
        value_cell.border = thin_border
    for column in (1, 2):
        sheet.cell(5, column).number_format = '#,##0.00;[Red](#,##0.00);-'
    for column in (4, 5):
        sheet.cell(5, column).number_format = '#,##0;[Red](#,##0);-'
    for column in (3, 6, 7):
        sheet.cell(5, column).number_format = '0.0%;[Red](0.0%);-'

    row = 8
    row = _write_section_title(sheet, row, "Monthly Performance", navy, white)
    monthly_headers = [
        "Month",
        f"Sales {previous_year}",
        f"Sales {year}",
        "Sales YoY",
        "Sales MoM",
        f"Qty {previous_year}",
        f"Qty {year}",
        "Qty YoY",
    ]
    row = _write_headers(sheet, row, monthly_headers, sky, navy)
    for item in report["monthly"]:
        values = [
            item["monthKey"],
            item["previousAmount"],
            item["currentAmount"],
            _percent(item["amountYoY"]),
            _percent(item["amountMoM"]),
            item["previousQty"],
            item["currentQty"],
            _percent(item["qtyYoY"]),
        ]
        _write_values(sheet, row, values, thin_border)
        for column in (2, 3):
            sheet.cell(row, column).number_format = '#,##0.00;[Red](#,##0.00);-'
        for column in (4, 5, 8):
            sheet.cell(row, column).number_format = '0.0%;[Red](0.0%);-'
        for column in (6, 7):
            sheet.cell(row, column).number_format = '#,##0;[Red](#,##0);-'
        row += 1

    row += 2
    row = _write_section_title(sheet, row, "Top 10 Branches", navy, white)
    ranking_headers = [
        "Rank",
        "Branch Code",
        "Branch",
        f"Sales {previous_year}",
        f"Sales {year}",
        "Sales YoY",
        f"Qty {previous_year}",
        f"Qty {year}",
        "Qty YoY",
    ]
    row = _write_headers(sheet, row, ranking_headers, sky, navy)
    for rank, item in enumerate(report["topBranches"], 1):
        values = [
            rank,
            item["branchCode"],
            item["displayName"],
            item["previousAmount"],
            item["currentAmount"],
            _percent(item["amountYoY"]),
            item["previousQty"],
            item["currentQty"],
            _percent(item["qtyYoY"]),
        ]
        _write_ranking_row(sheet, row, values, thin_border)
        row += 1

    row += 2
    row = _write_section_title(sheet, row, "Top 15 SKU", navy, white)
    sku_headers = [
        "Rank",
        "SKU",
        "Description",
        f"Sales {previous_year}",
        f"Sales {year}",
        "Sales YoY",
        f"Qty {previous_year}",
        f"Qty {year}",
        "Qty YoY",
    ]
    row = _write_headers(sheet, row, sku_headers, sky, navy)
    for rank, item in enumerate(report["topSkus"], 1):
        values = [
            rank,
            item["sku"],
            item["description"],
            item["previousAmount"],
            item["currentAmount"],
            _percent(item["amountYoY"]),
            item["previousQty"],
            item["currentQty"],
            _percent(item["qtyYoY"]),
        ]
        _write_ranking_row(sheet, row, values, thin_border)
        row += 1

    widths = [10, 18, 42, 18, 18, 14, 16, 16, 14]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    sheet.freeze_panes = "A10"
    sheet.auto_filter.ref = f"A9:H{9 + len(report['monthly'])}"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _write_section_title(sheet, row: int, title: str, fill: str, text: str) -> int:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    cell = sheet.cell(row, 1, title)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=text, bold=True, size=11)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 23
    return row + 1


def _write_headers(sheet, row: int, headers: list[str], fill: str, text: str) -> int:
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row, column, header)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color=text, bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 28
    return row + 1


def _write_values(sheet, row: int, values: list, border: Border) -> None:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.border = border
        cell.alignment = Alignment(vertical="center")


def _write_ranking_row(sheet, row: int, values: list, border: Border) -> None:
    _write_values(sheet, row, values, border)
    sheet.cell(row, 2).number_format = "@"
    for column in (4, 5):
        sheet.cell(row, column).number_format = '#,##0.00;[Red](#,##0.00);-'
    for column in (6, 9):
        sheet.cell(row, column).number_format = '0.0%;[Red](0.0%);-'
    for column in (7, 8):
        sheet.cell(row, column).number_format = '#,##0;[Red](#,##0);-'
