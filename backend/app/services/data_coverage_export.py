from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.local_time import as_bangkok, as_bangkok_excel, bangkok_now


@dataclass(frozen=True)
class CoverageBatch:
    batch_id: int
    data_date: date
    status: str
    branch_count: int
    item_count: int
    row_count: int
    source_filename: str
    finished_at: datetime | None


def data_coverage_filename(
    mt_code: str,
    year: int,
    current_time: datetime | None = None,
) -> str:
    timestamp = as_bangkok(current_time or bangkok_now()).strftime("%Y%m%d_%H%M%S")
    return f"{mt_code}_Data_Coverage_{year}_{timestamp}.xlsx"


def _dates_in_range(start: date, end: date) -> list[date]:
    return [start + timedelta(days=offset) for offset in range((end - start).days + 1)]


def _thai_day_name(value: date) -> str:
    return (
        "วันจันทร์",
        "วันอังคาร",
        "วันพุธ",
        "วันพฤหัสบดี",
        "วันศุกร์",
        "วันเสาร์",
        "วันอาทิตย์",
    )[value.weekday()]


def build_data_coverage_workbook(
    *,
    mt_code: str,
    mt_name: str,
    year: int,
    end_date: date,
    batches: list[CoverageBatch],
    current_time: datetime | None = None,
) -> bytes:
    start_date = date(year, 1, 1)
    dates = _dates_in_range(start_date, end_date)
    batch_by_date = {batch.data_date: batch for batch in batches}
    missing_dates = [value for value in dates if value not in batch_by_date]

    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "สรุป"
    detail = workbook.create_sheet("รายละเอียดรายวัน")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:D1")
    summary["A1"] = "รายงานความครบถ้วนของข้อมูล"
    summary["A1"].fill = PatternFill("solid", fgColor="102A43")
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30

    summary_rows = (
        ("Modern Trade", f"{mt_name} ({mt_code})"),
        ("ปีที่ตรวจสอบ", year),
        ("ตรวจสอบถึงวันที่", end_date),
        ("จำนวนวันที่ต้องมีข้อมูล", None),
        ("วันที่มีข้อมูล", None),
        ("วันที่ขาดข้อมูล", None),
        ("ความครบถ้วน", None),
        ("วันที่สร้างรายงาน", as_bangkok_excel(current_time or bangkok_now())),
    )
    for row_number, (label, value) in enumerate(summary_rows, start=3):
        summary.cell(row_number, 1, label)
        summary.cell(row_number, 2, value)
        summary.cell(row_number, 1).font = Font(bold=True, color="334E68")
        summary.cell(row_number, 1).fill = PatternFill("solid", fgColor="E8F1F1")

    last_detail_row = len(dates) + 1
    summary["B6"] = f"=COUNTA('รายละเอียดรายวัน'!A2:A{last_detail_row})"
    summary["B7"] = f'=COUNTIF(\'รายละเอียดรายวัน\'!D2:D{last_detail_row},"มีข้อมูล")'
    summary["B8"] = f'=COUNTIF(\'รายละเอียดรายวัน\'!D2:D{last_detail_row},"ขาดข้อมูล")'
    summary["B9"] = "=IFERROR(B7/B6,0)"
    summary["B5"].number_format = "dd/mm/yyyy"
    summary["B9"].number_format = "0.0%"
    summary["B10"].number_format = "dd/mm/yyyy hh:mm:ss"
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 32
    summary.column_dimensions["C"].width = 18
    summary.column_dimensions["D"].width = 18

    summary["A12"] = "วันที่ขาดข้อมูล"
    summary["A12"].fill = PatternFill("solid", fgColor="C6283D")
    summary["A12"].font = Font(color="FFFFFF", bold=True)
    for row_number, missing_date in enumerate(missing_dates, start=13):
        summary.cell(row_number, 1, missing_date)
        summary.cell(row_number, 1).number_format = "dd/mm/yyyy"
        summary.cell(row_number, 2, _thai_day_name(missing_date))
        summary.cell(row_number, 3, "วันหยุดสุดสัปดาห์" if missing_date.weekday() >= 5 else "วันทำงาน")
        for column in range(1, 4):
            summary.cell(row_number, column).fill = PatternFill("solid", fgColor="FDE2E5")
    summary.freeze_panes = "A3"

    headers = (
        "วันที่",
        "วัน",
        "ประเภทวัน",
        "สถานะข้อมูล",
        "Batch ID",
        "Branch",
        "Item",
        "จำนวนแถว",
        "ผลการ Import",
        "ไฟล์ต้นทาง",
        "Import เมื่อ",
    )
    detail.append(headers)
    for cell in detail[1]:
        cell.fill = PatternFill("solid", fgColor="0B756E")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    detail.row_dimensions[1].height = 24
    detail.freeze_panes = "A2"
    detail.sheet_view.showGridLines = False

    weekend_fill = PatternFill("solid", fgColor="EEF2F3")
    missing_fill = PatternFill("solid", fgColor="FDE2E5")
    weekend_marker_fill = PatternFill("solid", fgColor="D9E1E3")
    for row_number, data_date in enumerate(dates, start=2):
        batch = batch_by_date.get(data_date)
        is_weekend = data_date.weekday() >= 5
        import_result = ""
        if batch:
            import_result = (
                "สำเร็จพร้อมคำเตือน"
                if batch.status == "imported_with_warnings"
                else "สำเร็จ"
            )
        finished_at = as_bangkok_excel(batch.finished_at) if batch and batch.finished_at else None
        values = (
            data_date,
            _thai_day_name(data_date),
            "วันหยุดสุดสัปดาห์" if is_weekend else "วันทำงาน",
            "มีข้อมูล" if batch else "ขาดข้อมูล",
            batch.batch_id if batch else None,
            batch.branch_count if batch else None,
            batch.item_count if batch else None,
            batch.row_count if batch else None,
            import_result,
            batch.source_filename if batch else "",
            finished_at,
        )
        for column, value in enumerate(values, start=1):
            detail.cell(row_number, column, value)
        detail.cell(row_number, 1).number_format = "dd/mm/yyyy"
        detail.cell(row_number, 11).number_format = "dd/mm/yyyy hh:mm:ss"
        row_fill = missing_fill if not batch else weekend_fill if is_weekend else None
        if row_fill:
            for column in range(1, len(headers) + 1):
                detail.cell(row_number, column).fill = row_fill
        if is_weekend and not batch:
            for column in (2, 3):
                detail.cell(row_number, column).fill = weekend_marker_fill
        if not batch:
            detail.cell(row_number, 4).font = Font(color="C6283D", bold=True)

    widths = (14, 15, 22, 16, 12, 12, 12, 14, 22, 42, 22)
    for column, width in enumerate(widths, start=1):
        detail.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    table = Table(displayName=f"{mt_code}DataCoverage", ref=f"A1:K{last_detail_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    detail.add_table(table)
    detail.auto_filter.ref = f"A1:K{last_detail_row}"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
