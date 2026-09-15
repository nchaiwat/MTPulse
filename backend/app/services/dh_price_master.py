from __future__ import annotations

import hashlib
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.importers.dh import STORAGE_SCALE
from app.models import DhEffectivePrice, ModernTrade
from app.services.dh_pricing import DhPrice

SHEET_NAME = "DH Price Master"
HEADERS = ("SKU", "Price Ex VAT", "Effective From", "Effective To")


class DhPriceMasterError(ValueError):
    pass


@dataclass(frozen=True)
class DhPriceCandidate:
    source_sku: str
    unit_price_ex_vat: Decimal
    effective_from: date
    effective_to: date | None


@dataclass(frozen=True)
class ParsedDhPriceWorkbook:
    row_count: int
    candidates: tuple[DhPriceCandidate, ...]
    errors: tuple[str, ...]


@dataclass(frozen=True)
class DhPriceMasterPreview:
    row_count: int
    candidate_count: int
    inserted: int
    updated: int
    unchanged: int
    source_checksum_sha256: str
    errors: tuple[str, ...]


def build_dh_price_template() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    sheet.append(("00123", None, None, None))
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 24
    widths = (22, 20, 20, 20)
    header_fill = PatternFill("solid", fgColor="0B756E")
    for column, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    _set_text(sheet["A2"], "00123")
    sheet["B2"].number_format = "#,##0.0000"
    sheet["C2"].number_format = "yyyy-mm-dd"
    sheet["D2"].number_format = "yyyy-mm-dd"
    table = Table(displayName="DHPriceMaster", ref="A1:D2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    instructions = workbook.create_sheet("วิธีใช้งาน")
    instructions.column_dimensions["A"].width = 110
    notes = (
        "1. แทนที่แถวตัวอย่างด้วยราคาจริงของ DoHome",
        "2. SKU เป็นข้อความ ห้ามตัดเลขศูนย์นำหน้าหรือเติมเลขศูนย์เอง",
        "3. Price Ex VAT ต้องมากกว่า 0",
        "4. Effective From ต้องมีค่า; Effective To เว้นว่างได้",
        "5. ช่วงวันที่ของ SKU เดียวกันห้ามซ้อนกัน",
        "6. วันที่รองรับ Excel Date หรือข้อความรูปแบบ YYYY-MM-DD",
    )
    for row_number, note in enumerate(notes, start=1):
        _set_text(instructions.cell(row=row_number, column=1), note)
    instructions["A1"].font = Font(bold=True, color="0B756E")

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse_dh_price_workbook(content: bytes) -> ParsedDhPriceWorkbook:
    try:
        workbook = openpyxl.load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise DhPriceMasterError("ไฟล์ Excel ไม่ถูกต้องหรือเปิดอ่านไม่ได้") from exc

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise DhPriceMasterError(f"ไม่พบ Sheet '{SHEET_NAME}'")
        sheet = workbook[SHEET_NAME]
        first_row = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        actual_headers = tuple(_text(value) for value in first_row)
        if actual_headers != HEADERS:
            raise DhPriceMasterError(
                "Column ต้องเรียงตามลำดับ: " + ", ".join(HEADERS)
            )

        row_count = 0
        errors: list[str] = []
        candidates_by_key: dict[
            tuple[str, date],
            list[tuple[int, DhPriceCandidate]],
        ] = {}
        for excel_row, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not any(value not in (None, "") for value in values):
                continue
            row_count += 1
            candidate = _parse_candidate(values, excel_row, errors)
            if candidate is None:
                continue
            key = (candidate.source_sku, candidate.effective_from)
            candidates_by_key.setdefault(key, []).append((excel_row, candidate))

        for (sku, effective_from), rows in candidates_by_key.items():
            if len(rows) > 1:
                row_labels = ", ".join(str(row_number) for row_number, _ in rows)
                errors.append(
                    f"SKU {sku} วันที่ {effective_from.isoformat()} ซ้ำที่แถว "
                    f"{row_labels}"
                )

        if row_count == 0:
            errors.append("ไม่พบข้อมูลราคาในไฟล์")
        if errors:
            return ParsedDhPriceWorkbook(
                row_count=row_count,
                candidates=(),
                errors=tuple(errors),
            )
        candidates = tuple(
            rows[0][1]
            for _, rows in sorted(candidates_by_key.items())
        )
        return ParsedDhPriceWorkbook(
            row_count=row_count,
            candidates=candidates,
            errors=(),
        )
    finally:
        workbook.close()


def preview_dh_price_master(
    session: Session,
    content: bytes,
) -> DhPriceMasterPreview:
    parsed = parse_dh_price_workbook(content)
    checksum = hashlib.sha256(content).hexdigest()
    if parsed.errors:
        return DhPriceMasterPreview(
            row_count=parsed.row_count,
            candidate_count=0,
            inserted=0,
            updated=0,
            unchanged=0,
            source_checksum_sha256=checksum,
            errors=parsed.errors,
        )

    modern_trade = session.scalar(
        select(ModernTrade).where(ModernTrade.code == "DH")
    )
    if modern_trade is None:
        raise DhPriceMasterError("ยังไม่มี Modern Trade รหัส DH ในฐานข้อมูล")
    existing_rows = session.scalars(
        select(DhEffectivePrice)
        .where(DhEffectivePrice.modern_trade_id == modern_trade.id)
        .order_by(DhEffectivePrice.source_sku, DhEffectivePrice.effective_from)
    ).all()
    existing_by_key = {
        (row.source_sku, row.effective_from): row
        for row in existing_rows
    }
    proposed: dict[tuple[str, date], DhPriceCandidate] = {
        (row.source_sku, row.effective_from): DhPriceCandidate(
            source_sku=row.source_sku,
            unit_price_ex_vat=row.unit_price_ex_vat.quantize(STORAGE_SCALE),
            effective_from=row.effective_from,
            effective_to=row.effective_to,
        )
        for row in existing_rows
    }
    inserted = updated = unchanged = 0
    for candidate in parsed.candidates:
        key = (candidate.source_sku, candidate.effective_from)
        existing = existing_by_key.get(key)
        if existing is None:
            inserted += 1
        elif (
            existing.unit_price_ex_vat.quantize(STORAGE_SCALE)
            == candidate.unit_price_ex_vat
            and existing.effective_to == candidate.effective_to
        ):
            unchanged += 1
        else:
            updated += 1
        proposed[key] = candidate

    errors = _overlap_errors(proposed.values())
    return DhPriceMasterPreview(
        row_count=parsed.row_count,
        candidate_count=len(parsed.candidates),
        inserted=inserted,
        updated=updated,
        unchanged=unchanged,
        source_checksum_sha256=checksum,
        errors=errors,
    )


def load_effective_dh_prices(
    session: Session,
    sales_date: date,
    source_skus: set[str],
) -> tuple[DhPrice, ...]:
    if not source_skus:
        return ()
    modern_trade = session.scalar(
        select(ModernTrade).where(ModernTrade.code == "DH")
    )
    if modern_trade is None:
        raise DhPriceMasterError("ยังไม่มี Modern Trade รหัส DH ในฐานข้อมูล")
    rows = session.scalars(
        select(DhEffectivePrice)
        .where(
            DhEffectivePrice.modern_trade_id == modern_trade.id,
            DhEffectivePrice.source_sku.in_(sorted(source_skus)),
            DhEffectivePrice.effective_from <= sales_date,
            (
                DhEffectivePrice.effective_to.is_(None)
                | (DhEffectivePrice.effective_to >= sales_date)
            ),
        )
        .order_by(DhEffectivePrice.source_sku, DhEffectivePrice.effective_from)
    ).all()
    return tuple(
        DhPrice(
            source_sku=row.source_sku,
            unit_price_ex_vat=row.unit_price_ex_vat,
            effective_from=row.effective_from,
            effective_to=row.effective_to,
        )
        for row in rows
    )


def _parse_candidate(
    values: tuple[object, ...],
    excel_row: int,
    errors: list[str],
) -> DhPriceCandidate | None:
    sku = _identifier(values[0] if values else None)
    if not sku:
        errors.append(f"แถว {excel_row}: ไม่มี SKU")
    elif len(sku) > 50:
        errors.append(f"แถว {excel_row}: SKU ยาวเกิน 50 ตัวอักษร")

    unit_price = _price(
        values[1] if len(values) > 1 else None,
        excel_row,
        errors,
    )
    effective_from = _date_value(
        values[2] if len(values) > 2 else None,
        excel_row,
        "Effective From",
        required=True,
        errors=errors,
    )
    effective_to = _date_value(
        values[3] if len(values) > 3 else None,
        excel_row,
        "Effective To",
        required=False,
        errors=errors,
    )
    if (
        effective_from is not None
        and effective_to is not None
        and effective_to < effective_from
    ):
        errors.append(f"แถว {excel_row}: ช่วงวันที่ไม่ถูกต้อง")
    if (
        not sku
        or len(sku) > 50
        or unit_price is None
        or effective_from is None
        or (effective_to is not None and effective_to < effective_from)
    ):
        return None
    return DhPriceCandidate(
        source_sku=sku,
        unit_price_ex_vat=unit_price,
        effective_from=effective_from,
        effective_to=effective_to,
    )


def _overlap_errors(
    candidates: Iterable[DhPriceCandidate],
) -> tuple[str, ...]:
    by_sku: dict[str, list[DhPriceCandidate]] = {}
    for candidate in candidates:
        by_sku.setdefault(candidate.source_sku, []).append(candidate)
    errors = []
    for sku, rows in sorted(by_sku.items()):
        ordered = sorted(rows, key=lambda row: row.effective_from)
        for previous, current in zip(ordered, ordered[1:], strict=False):
            if (
                previous.effective_to is None
                or current.effective_from <= previous.effective_to
            ):
                errors.append(
                    f"ช่วงราคาของ SKU {sku} ซ้อนกันระหว่าง "
                    f"{previous.effective_from.isoformat()} และ "
                    f"{current.effective_from.isoformat()}"
                )
    return tuple(errors)


def _identifier(value: object) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _price(
    value: object,
    excel_row: int,
    errors: list[str],
) -> Decimal | None:
    try:
        unit_price = Decimal(str(value).replace(",", ""))
    except (InvalidOperation, AttributeError):
        errors.append(f"แถว {excel_row}: Price Ex VAT ต้องเป็นตัวเลข")
        return None
    if not unit_price.is_finite() or unit_price <= 0:
        errors.append(f"แถว {excel_row}: Price Ex VAT ต้องมากกว่า 0")
        return None
    if unit_price.adjusted() > 15:
        errors.append(
            f"แถว {excel_row}: Price Ex VAT มีจำนวนหลักเกินขนาดที่รองรับ"
        )
        return None
    try:
        quantized = unit_price.quantize(STORAGE_SCALE)
    except InvalidOperation:
        errors.append(
            f"แถว {excel_row}: Price Ex VAT มีจำนวนหลักเกินขนาดที่รองรับ"
        )
        return None
    if quantized <= 0:
        errors.append(f"แถว {excel_row}: Price Ex VAT เล็กเกิน precision ที่รองรับ")
        return None
    return quantized


def _date_value(
    value: object,
    excel_row: int,
    column: str,
    *,
    required: bool,
    errors: list[str],
) -> date | None:
    if value in (None, ""):
        if required:
            errors.append(f"แถว {excel_row}: ไม่มี {column}")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    errors.append(f"แถว {excel_row}: {column} ต้องเป็นวันที่ YYYY-MM-DD")
    return None


def _text(value: object) -> str:
    return "" if value is None else str(value).replace("\u00a0", " ").strip()


def _set_text(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"
    cell.number_format = "@"
