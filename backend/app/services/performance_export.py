from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

from app.local_time import as_bangkok, bangkok_now

METRIC_LABELS = {
    "amount": "Amount",
    "qty": "Qty",
    "stockOh": "Stock On Hand",
    "stockOnOrder": "Stock On Order",
    "stockValue": "Stock Value (Source)",
}

def _display_date(value: str) -> str:
    return datetime.fromisoformat(value).strftime("%d/%m/%Y")


def _display_period(value: str) -> str:
    if len(value) == 7:
        return datetime.strptime(value, "%Y-%m").strftime("%b %Y")
    return _display_date(value)



def performance_export_filename(
    mode: str,
    metric: str,
    grain: str,
    current_time: datetime | None = None,
    sales_basis: str = "net",
    mt_code: str = "TWD",
) -> str:
    timestamp = as_bangkok(current_time or bangkok_now()).strftime("%Y%m%d_%H%M%S")
    view = (
        "Branch"
        if grain in {"branch_month", "branch_range", "day"}
        else "Month"
        if grain == "month"
        else "Date"
    )
    basis = "_Gross" if mode == "sales" and sales_basis == "gross" else ""
    metric_label = METRIC_LABELS[metric].replace(" ", "")
    return f"{mt_code}_{mode.title()}{basis}_{metric_label}_{view}_{timestamp}.xlsx"


def build_performance_workbook(
    report: dict,
    *,
    mode: str,
    metric: str,
    grain: str,
    show_descriptions: bool,
    branch_id: str | None = None,
    branch_ids: list[str] | None = None,
    sales_basis: str = "net",
    mt_code: str = "TWD",
) -> bytes:
    dimension = (
        "branch"
        if grain in {"branch_month", "branch_range", "day"}
        else "month"
        if grain == "month"
        else "date"
    )
    dates = list(report["dates"])
    selected_dates = (
        [dates[-1]] if mode == "inventory" and dimension == "branch" and dates else dates
    )
    selected_branch_ids = set(branch_ids or ([branch_id] if branch_id else []))
    branches = [
        branch
        for branch in report["branches"]
        if not selected_branch_ids or branch["id"] in selected_branch_ids
    ]
    dimension_keys = (
        [branch["id"] for branch in branches] if dimension == "branch" else selected_dates
    )
    branch_names = {branch["id"]: branch["name"] for branch in branches}

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.sheet_view.showGridLines = False

    identity_headers = ["Sho", "Pro", f"{mt_code} SKU"]
    if show_descriptions:
        identity_headers.append(f"{mt_code} Description")
    identity_headers.append("WA Item")
    if show_descriptions:
        identity_headers.append("WA Description")
    show_turnover = mode == "inventory"
    if show_turnover:
        identity_headers.extend(["TOM", "TOD"])
    headers = (
        identity_headers
        + ["TOTAL"]
        + [
            (
                f"{key} - {branch_names[key]}"
                if dimension == "branch"
                else _display_date(key) if dimension == "date" else key
            )
            for key in dimension_keys
        ]
    )
    total_column = len(identity_headers) + 1
    data_start_row = 6
    data_end_row = data_start_row + len(report["items"]) - 1

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    basis_label = (
        "Gross Sale Out" if mode == "sales" and sales_basis == "gross" else "Net Sales"
    )
    title_basis = f" ({basis_label})" if mode == "sales" else ""
    sheet.cell(
        1,
        1,
        f"{mt_code} {mode.title()}{title_basis} by {dimension.title()} — {METRIC_LABELS[metric]}",
    )
    selected_ranges = report.get("selectedDateRanges") or []
    period = report.get("selectedMonth") or (
        ", ".join(
            _display_date(value["from"])
            if value["from"] == value["to"]
            else f"{_display_date(value['from'])} – {_display_date(value['to'])}"
            for value in selected_ranges
        )
        if selected_ranges
        else " – ".join(_display_period(value) for value in selected_dates)
        if selected_dates
        else "ไม่มีข้อมูล"
    )
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.cell(2, 1, f"ช่วงข้อมูล: {period} | Export: {bangkok_now():%d/%m/%Y %H:%M:%S}")
    sheet.cell(4, 1, "AVG" if show_turnover else "SUM")
    for column, header in enumerate(headers, start=1):
        sheet.cell(5, column, header)

    number_format = (
        '#,##0.00;[Red]-#,##0.00;-""'
        if metric in {"amount", "stockValue"}
        else '#,##0;[Red]-#,##0;-""'
    )
    for row_number, item in enumerate(report["items"], start=data_start_row):
        values = [
            "✓" if item.get("isSho") else "",
            "✓" if item.get("isPro") else "",
            item["sku"],
        ]
        if show_descriptions:
            values.append(item["twdDescription"])
        values.append(item["waItem"] or "")
        if show_descriptions:
            values.append(item["waDescription"] or "")
        if show_turnover:
            values.extend([item.get("tom"), item.get("tod")])
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            if column in {3, 5 if show_descriptions else 4}:
                cell.data_type = "s"
                cell.number_format = "@"
        if show_turnover:
            turnover_start = len(identity_headers) - 1
            sheet.cell(row_number, turnover_start).number_format = '#,##0.00;[Red]-#,##0.00;-""'
            sheet.cell(row_number, turnover_start + 1).number_format = '#,##0.00;[Red]-#,##0.00;-""'

        points = [
            point
            for point in item["points"]
            if point["date"] in selected_dates
            and (
                not selected_branch_ids
                or point["branchId"] == "all"
                or point["branchId"] in selected_branch_ids
            )
        ]
        totals = {key: 0.0 for key in dimension_keys}
        for point in points:
            key = point["branchId"] if dimension == "branch" else point["date"]
            if key in totals:
                totals[key] += point[metric]
        first_value_column = total_column + 1
        last_value_column = total_column + len(dimension_keys)
        total_cell = sheet.cell(row_number, total_column)
        if dimension_keys:
            first_letter = openpyxl.utils.get_column_letter(first_value_column)
            last_letter = openpyxl.utils.get_column_letter(last_value_column)
            total_cell.value = (
                f"=SUM({first_letter}{row_number}:{last_letter}{row_number})"
            )
        else:
            total_cell.value = 0
        total_cell.number_format = number_format
        for offset, key in enumerate(dimension_keys, start=first_value_column):
            cell = sheet.cell(row_number, offset, totals[key])
            cell.number_format = number_format

    if report["items"]:
        for column in range(total_column, len(headers) + 1):
            letter = openpyxl.utils.get_column_letter(column)
            cell = sheet.cell(4, column, f"=SUM({letter}{data_start_row}:{letter}{data_end_row})")
            cell.number_format = number_format
    else:
        sheet.cell(4, total_column, 0).number_format = number_format
    if show_turnover:
        turnover_start = len(identity_headers) - 1
        inventory_summary = report.get("inventorySummary") or {}
        sheet.cell(4, turnover_start, inventory_summary.get("averageTom"))
        sheet.cell(4, turnover_start + 1, inventory_summary.get("averageTod"))
        sheet.cell(4, turnover_start).number_format = '#,##0.00;[Red]-#,##0.00;-""'
        sheet.cell(4, turnover_start + 1).number_format = '#,##0.00;[Red]-#,##0.00;-""'

    title_fill = PatternFill("solid", fgColor="102A43")
    header_fill = PatternFill("solid", fgColor="0B756E")
    sum_fill = PatternFill("solid", fgColor="D8F1EE")
    for cell in sheet[1]:
        cell.fill = title_fill
        cell.font = Font(color="FFFFFF", bold=True, size=14)
    for cell in sheet[5]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet["A5"].fill = PatternFill("solid", fgColor="D97706")
    sheet["B5"].fill = PatternFill("solid", fgColor="7C3AED")
    for cell in sheet[4]:
        cell.fill = sum_fill
        cell.font = Font(color="102A43", bold=True)
    sheet["A2"].font = Font(color="526D82", italic=True, size=9)
    sheet.row_dimensions[1].height = 27
    sheet.row_dimensions[5].height = 40
    sheet.freeze_panes = sheet.cell(data_start_row, total_column + 1)
    sheet.auto_filter.ref = (
        f"A5:{openpyxl.utils.get_column_letter(len(headers))}{max(5, data_end_row)}"
    )

    widths = [8, 8, 14]
    if show_descriptions:
        widths.append(42)
    widths.append(24)
    if show_descriptions:
        widths.append(48)
    if show_turnover:
        widths.extend([12, 12])
    widths.extend([16] + [15] * len(dimension_keys))
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=data_start_row, max_row=max(data_start_row, data_end_row)):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
    for row_number, item in enumerate(report["items"], start=data_start_row):
        is_sho = bool(item.get("isSho"))
        is_pro = bool(item.get("isPro"))
        row_color = (
            "F6E8EC"
            if is_sho and is_pro
            else "FFF5DF"
            if is_sho
            else "F3E8FF"
            if is_pro
            else None
        )
        if row_color:
            for column in range(1, len(headers) + 1):
                sheet.cell(row_number, column).fill = PatternFill("solid", fgColor=row_color)
        if is_sho:
            sheet.cell(row_number, 1).fill = PatternFill("solid", fgColor="D97706")
            sheet.cell(row_number, 1).font = Font(color="FFFFFF", bold=True)
        if is_pro:
            sheet.cell(row_number, 2).fill = PatternFill("solid", fgColor="7C3AED")
            sheet.cell(row_number, 2).font = Font(color="FFFFFF", bold=True)

    if report["items"] and dimension_keys:
        value_range = (
            f"{openpyxl.utils.get_column_letter(total_column + 1)}{data_start_row}:"
            f"{openpyxl.utils.get_column_letter(len(headers))}{data_end_row}"
        )
        sheet.conditional_formatting.add(
            value_range,
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFF",
                mid_type="percentile",
                mid_value=50,
                mid_color="D8F1EE",
                end_type="max",
                end_color="77C8BD",
            ),
        )
        sheet.conditional_formatting.add(
            value_range,
            CellIsRule(
                operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FDE2E5")
            ),
        )

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
