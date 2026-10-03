# ruff: noqa: F811
import json
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from sqlalchemy import select
from test_ciam_auth import account, auth_client, sign_in  # noqa: F401

from app.models import (
    AssortmentBase,
    AssortmentForecast,
    AssortmentMember,
    ImportBatch,
    ItemMapping,
    ModernTrade,
    SalesInventoryFact,
    TransactionLog,
)
from app.services import assortment
from app.services.monthly_sales_summary import refresh_monthly_sales_summary


def test_assortment_viewer_cannot_mutate(auth_client):
    client, db = auth_client
    sign_in(client, db, account(db, role="viewer"))
    assert client.get("/api/assortment/catalog").status_code == 200
    assert client.post("/api/assortment/bases", json={"description": "A"}).status_code == 403


def test_assortment_schema_exists():
    assert AssortmentBase.__tablename__ == "assortment_bases"


def seed(db, code="TWD"):
    db.add(ModernTrade(id=1, code=code, name=code, show_unmatched_branches=True))
    for index, wa in enumerate(["WA-01", "WA-02", "WA-unconfirmed"], 1):
        db.add(
            ItemMapping(
                id=index,
                modern_trade_id=1,
                source_sku=str(index),
                wa_item_code=wa,
                wa_item_description="ประตู สีขาว",
                status="confirmed",
                effective_from=date(2025, 1, 1),
                changed_by="test",
            )
        )
    for day, grain in [(1, "daily"), (2, "daily"), (3, "rolling_30d")]:
        db.add(
            ImportBatch(
                id=day,
                modern_trade_id=1,
                data_date=date(2026, 1, day),
                sales_grain=grain,
                status="imported",
                source_path="test",
                source_filename="test.xls",
                checksum_sha256=str(day),
                row_count=3,
                store_count=1,
                sku_count=3,
                negative_row_count=0,
                source_amount=0,
                amount=0,
                sales_qty=0,
                stock_on_hand=0,
                reported_stock_on_hand=0,
                stock_on_order=0,
            )
        )
        for i in range(1, 4):
            db.add(
                SalesInventoryFact(
                    id=day * 10 + i,
                    modern_trade_id=1,
                    batch_id=day,
                    data_date=date(2026, 1, day),
                    sales_grain=grain,
                    source_sku=str(i),
                    source_branch_code="X",
                    source_branch_name="X",
                    source_amount=0,
                    amount=0,
                    sales_qty=100 if day == 3 else (2 if day == 1 else -1),
                    stock_on_hand=day * 10,
                    stock_on_order=0,
                )
            )
    db.add(AssortmentBase(id="b", description="ประตู", attributes=json.dumps([""] * 7), version=1))
    db.add_all([AssortmentMember(wa_item_code=wa, base_id="b") for wa in ["WA-01", "WA-02"]])
    db.commit()
    refresh_monthly_sales_summary(db, 1, date(2026, 1, 1))
    db.commit()


@pytest.mark.parametrize("code", ["TWD", "HP", "MH", "HH", "GH", "DH", "TA"])
def test_report_reuses_mt_rules_and_excludes_unconfirmed(auth_client, code):
    _, db = auth_client
    seed(db, code)
    result = assortment.report(db, 2027, mt=code)
    c = result["items"][0]["cells"][code]
    assert c["years"]["2026"]["total"] == 2
    assert c["years"]["2025"]["total"] is None
    assert c["forecastTotal"] is None and c["growth"] is None
    assert result["coverage"][code]["2026"]["days"] == 2
    assert result["coverage"][code]["2026"]["months"][0] == {"days": 2, "expected": 31}
    assert (
        assortment.report(db, 2027, mt=code, basis="gross")["items"][0]["cells"][code]["years"][
            "2026"
        ]["total"]
        == 4
    )
    inv = assortment.report(db, 2027, mt=code, mode="inventory")["items"][0]["cells"][code]
    assert inv["years"]["2026"]["total"] == 60
    assert inv["years"]["2026"]["snapshot"] == "2026-01-03"


def test_forecast_version_clone_export_and_audit(auth_client):
    client, db = auth_client
    seed(db)
    sign_in(client, db, account(db))
    plan = client.post("/api/assortment/plans", json={"year": 2027, "name": "Plan"}).json()
    path = f"/api/assortment/plans/{plan['id']}/forecast/b/TWD"
    values = [100, 0] + [None] * 10
    response = client.put(path, json={"version": 1, "months": values})
    assert response.status_code == 200, response.text
    assert client.put(path, json={"version": 1, "months": [9] * 12}).status_code == 409
    assert client.put(path, json={"version": 2, "months": [-1] * 12}).status_code == 422
    copy = client.post(
        "/api/assortment/plans", json={"year": 2027, "name": "Copy", "copy_id": plan["id"]}
    ).json()
    assert (
        client.put(
            f"/api/assortment/plans/{copy['id']}/forecast/b/TWD",
            json={"version": 1, "months": [9] * 12},
        ).status_code
        == 200
    )
    original = json.loads(db.get(AssortmentForecast, (plan["id"], "b", "TWD")).months)
    assert original[:3] == ["100", "0", None]
    assert (
        client.put(f"/api/assortment/plans/{copy['id']}/primary", json={"version": 2}).status_code
        == 200
    )
    response = client.get(
        f"/api/assortment/report?year=2027&plan_id={plan['id']}&mt=TWD&download=true"
    )
    assert response.status_code == 200, response.text
    workbook = openpyxl.load_workbook(BytesIO(response.content))
    assert workbook["Forecast"].cell(2, 12).value == 100
    assert workbook["Forecast"].cell(2, 13).value == 0
    assert workbook["Forecast"].cell(2, 14).value is None
    assert workbook["Assortment"].cell(4, 12).value == 100
    assert (
        db.scalar(select(TransactionLog).where(TransactionLog.action == "forecast")).triggered_by
        == "user:tester"
    )


def test_mapping_conflict_stale_and_admin_only(auth_client):
    client, db = auth_client
    seed(db)
    user = account(db)
    sign_in(client, db, user)
    payload = {"description": "Other", "members": ["WA-01"]}
    assert client.post("/api/assortment/bases", json=payload).status_code == 409
    assert (
        client.post(
            "/api/assortment/bases", json={"description": "Other", "members": ["fake"]}
        ).status_code
        == 422
    )
    payload = {"description": "Changed", "members": ["WA-01"], "version": 1}
    assert client.put("/api/assortment/bases/b", json=payload).status_code == 200
    assert client.put("/api/assortment/bases/b", json=payload).status_code == 409
    user.role = "operator"
    db.commit()
    assert client.put("/api/assortment/bases/b", json=payload).status_code == 403
    assert client.put("/api/assortment/display", json={"page_size": 50}).status_code == 403
    assert (
        client.post(
            "/api/assortment/plans", json={"name": "Operator plan", "year": 2027}
        ).status_code
        == 200
    )


def test_export_covers_all_pages_and_user_text_is_not_formula(auth_client):
    client, db = auth_client
    sign_in(client, db, account(db))
    db.add_all(
        [
            AssortmentBase(
                id=str(i),
                description="=SUM(A1)" if i == 0 else str(i),
                attributes=json.dumps([""] * 7),
                version=1,
            )
            for i in range(30)
        ]
    )
    db.commit()
    assert len(client.get("/api/assortment/report?mt=TWD").json()["items"]) == 25
    wb = openpyxl.load_workbook(
        BytesIO(client.get("/api/assortment/report?mt=TWD&download=true").content)
    )
    assert wb["Assortment"].max_row == 33
    found = [c for row in wb["Assortment"] for c in row if c.value == "=SUM(A1)"]
    assert found and all(c.data_type == "s" for c in found)


def test_create_base_requires_explicit_confirmed_membership(auth_client):
    client, db = auth_client
    seed(db)
    sign_in(client, db, account(db))
    response = client.post(
        "/api/assortment/bases", json={"description": "New base", "members": ["WA-unconfirmed"]}
    )
    assert response.status_code == 200, response.text
    created = response.json()
    assert created["version"] == 1
    assert db.get(AssortmentMember, "WA-unconfirmed").base_id == created["id"]
    report = client.get("/api/assortment/report?mt=TWD&search=New").json()
    assert report["items"][0]["cells"]["TWD"]["years"]["2026"]["total"] == 1
