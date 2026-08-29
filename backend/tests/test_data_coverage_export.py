from datetime import UTC, date, datetime
from io import BytesIO

import openpyxl

from app.services.data_coverage_export import (
    CoverageBatch,
    build_data_coverage_workbook,
    data_coverage_filename,
)


def test_data_coverage_workbook_lists_every_day_and_structural_counts() -> None:
    content = build_data_coverage_workbook(
        mt_code="TWD",
        mt_name="ไทวัสดุ",
        year=2025,
        end_date=date(2025, 12, 31),
        batches=[
            CoverageBatch(
                batch_id=9,
                data_date=date(2025, 1, 4),
                status="imported_with_warnings",
                branch_count=90,
                item_count=2043,
                row_count=12560,
                source_filename="source.xls",
                finished_at=datetime(2025, 1, 5, 8, 30),
            )
        ],
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    summary = workbook["สรุป"]
    detail = workbook["รายละเอียดรายวัน"]

    assert detail.max_row == 366
    assert detail["A2"].value == datetime(2025, 1, 1)
    assert detail["A5"].value == datetime(2025, 1, 4)
    assert detail["C5"].value == "วันหยุดสุดสัปดาห์"
    assert detail["D5"].value == "มีข้อมูล"
    assert detail["E5"].value == 9
    assert detail["F5"].value == 90
    assert detail["G5"].value == 2043
    assert detail["H5"].value == 12560
    assert detail["I5"].value == "สำเร็จพร้อมคำเตือน"
    assert detail["J5"].value == "source.xls"
    assert detail["D6"].value == "ขาดข้อมูล"
    assert detail["C6"].value == "วันหยุดสุดสัปดาห์"
    assert detail["D6"].fill.fgColor.rgb == "00FDE2E5"
    assert detail["C6"].fill.fgColor.rgb == "00D9E1E3"
    assert summary["B6"].value == "=COUNTA('รายละเอียดรายวัน'!A2:A366)"
    assert summary["B7"].value == '=COUNTIF(\'รายละเอียดรายวัน\'!D2:D366,"มีข้อมูล")'
    assert summary["B8"].value == '=COUNTIF(\'รายละเอียดรายวัน\'!D2:D366,"ขาดข้อมูล")'
    assert summary["A13"].value == datetime(2025, 1, 1)
    workbook.close()


def test_data_coverage_filename_contains_mt_year_and_timestamp() -> None:
    filename = data_coverage_filename(
        "TWD",
        2026,
        datetime(2026, 8, 25, 14, 35, 10),
    )

    assert filename == "TWD_Data_Coverage_2026_20260825_143510.xlsx"


def test_data_coverage_times_are_exported_in_bangkok_time() -> None:
    content = build_data_coverage_workbook(
        mt_code="TWD",
        mt_name="ไทวัสดุ",
        year=2026,
        end_date=date(2026, 1, 1),
        batches=[
            CoverageBatch(
                batch_id=1,
                data_date=date(2026, 1, 1),
                status="imported",
                branch_count=90,
                item_count=113,
                row_count=1000,
                source_filename="source.xlsx",
                finished_at=datetime(2026, 1, 1, 1, 30, tzinfo=UTC),
            )
        ],
        current_time=datetime(2026, 8, 29, 4, 25, 37, tzinfo=UTC),
    )

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    assert workbook["สรุป"]["B10"].value == datetime(2026, 8, 29, 11, 25, 37)
    assert workbook["รายละเอียดรายวัน"]["K2"].value == datetime(2026, 1, 1, 8, 30)
    workbook.close()

    assert data_coverage_filename(
        "TWD",
        2026,
        datetime(2026, 8, 29, 17, 30, tzinfo=UTC),
    ) == "TWD_Data_Coverage_2026_20260830_003000.xlsx"
