from datetime import date
from pathlib import Path

from openpyxl import Workbook

from app.importers.dh import DhFormatError, extract_dh_pair, inspect_dh_workbook


def _save_sale(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["รหัสสินค้า", "ชื่อสินค้า", "ชื่อผู้ดูแลขาย", "บางนา-ตราด", "To go สาทร", "รวม"])
    sheet.append(["00123", "สินค้า A", "OWNER", 2, 0, 2])
    sheet.append(["00456", "สินค้า B", "OWNER", -1, 3, 2])
    sheet.append([None, None, None, 1_500.25, 3_000, 4])
    workbook.save(path)


def _save_stock(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "รหัสสาขา",
            "ชื่อสาขา",
            "รหัสสินค้า",
            "ชื่อสินค้า",
            "หน่วย",
            "ชื่อหน่วย",
            "สต็อกคงเหลือ",
            "จำนวนขาย",
        ]
    )
    sheet.append(["BN", "บางนา-ตราด", "00123", "สินค้า A", "EA", "ชิ้น", 10, 2])
    sheet.append(["BN", "บางนา-ตราด", "00456", "สินค้า B", "EA", "ชิ้น", 5, -1])
    workbook.save(path)


def test_dh_pair_keeps_real_sales_stock_and_batch_dates(tmp_path: Path) -> None:
    sale = tmp_path / "รายงานยอดขาย06-01-2025_06-01-2025.xlsx"
    stock = tmp_path / "รายงานสต็อคAll07-01-2025_07-01-2025.xlsx"
    _save_sale(sale)
    _save_stock(stock)

    extract = extract_dh_pair(stock, sale)

    assert extract.batch_date == date(2025, 1, 7)
    assert extract.sales_date == date(2025, 1, 6)
    assert extract.stock_date == date(2025, 1, 7)
    assert extract.sales_summary.sales_qty == 4
    assert extract.sales_summary.source_amount == 4_500.25
    assert extract.stock_summary.stock_on_hand == 15
    assert {row.branch_code for row in extract.sales_rows if row.branch_name == "บางนา-ตราด"} == {"BN"}
    assert len({row.branch_code for row in extract.sales_rows if row.branch_name == "To go สาทร"}) == 1
    assert extract.reconciliation_errors == ()


def test_dh_inspection_detects_kind_and_source_date(tmp_path: Path) -> None:
    sale = tmp_path / "รายงานยอดขาย06-01-2025_06-01-2025.xlsx"
    stock = tmp_path / "รายงานสต็อคAll07-01-2025_07-01-2025.xlsx"
    _save_sale(sale)
    _save_stock(stock)

    assert inspect_dh_workbook(sale) == ("sales", date(2025, 1, 6))
    assert inspect_dh_workbook(stock) == ("inventory", date(2025, 1, 7))


def test_dh_pair_rejects_non_consecutive_source_dates(tmp_path: Path) -> None:
    sale = tmp_path / "รายงานยอดขาย05-01-2025_05-01-2025.xlsx"
    stock = tmp_path / "รายงานสต็อคAll07-01-2025_07-01-2025.xlsx"
    _save_sale(sale)
    _save_stock(stock)

    try:
        extract_dh_pair(stock, sale)
    except DhFormatError as exc:
        assert "ก่อน Stock 1 วัน" in str(exc)
    else:
        raise AssertionError("expected a DH date validation error")


def test_dh_pair_blocks_when_sale_footer_does_not_reconcile(tmp_path: Path) -> None:
    sale = tmp_path / "รายงานยอดขาย06-01-2025_06-01-2025.xlsx"
    stock = tmp_path / "รายงานสต็อคAll07-01-2025_07-01-2025.xlsx"
    _save_sale(sale)
    _save_stock(stock)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["รหัสสินค้า", "ชื่อสินค้า", "ชื่อผู้ดูแลขาย", "บางนา-ตราด", "รวม"])
    sheet.append(["00123", "สินค้า A", "OWNER", 2, 2])
    sheet.append([None, None, None, 1_500, 3])
    workbook.save(sale)

    extract = extract_dh_pair(stock, sale)

    assert extract.reconciliation_errors == ("QTY รวมจากรายการ 2 ไม่ตรงกับ Total 3",)
