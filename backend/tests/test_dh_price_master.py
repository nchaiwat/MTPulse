from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from sqlalchemy import create_engine, inspect, select
from sqlalchemy.orm import Session

from app.database import Base
from app.models import DhEffectivePrice, ModernTrade
from app.services.dh_price_master import (
    DhPriceMasterError,
    build_dh_price_template,
    load_effective_dh_prices,
    parse_dh_price_workbook,
    preview_dh_price_master,
)

HEADERS = ("SKU", "Price Ex VAT", "Effective From", "Effective To")


def _workbook_bytes(rows: list[tuple[object, ...]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "DH Price Master"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _database() -> tuple[object, Session]:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    return engine, Session(engine)


def _price(
    row_id: int,
    modern_trade_id: int,
    sku: str,
    amount: int,
    effective_from: date,
    effective_to: date | None = None,
) -> DhEffectivePrice:
    return DhEffectivePrice(
        id=row_id,
        modern_trade_id=modern_trade_id,
        source_sku=sku,
        unit_price_ex_vat=Decimal(amount),
        effective_from=effective_from,
        effective_to=effective_to,
        source_filename="existing.xlsx",
        source_checksum_sha256="a" * 64,
        changed_by="test",
    )


def test_dh_price_template_preserves_sku_as_opaque_text() -> None:
    content = build_dh_price_template()

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    sheet = workbook["DH Price Master"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet["A2"].value == "00123"
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].number_format == "@"
    workbook.close()


def test_parse_dh_price_workbook_accepts_excel_and_iso_dates_without_padding_sku() -> None:
    parsed = parse_dh_price_workbook(
        _workbook_bytes(
            [
                ("00123", 100.25, date(2025, 1, 1), date(2025, 1, 31)),
                (123, 200, "2025-02-01", None),
            ]
        )
    )

    assert parsed.row_count == 2
    assert parsed.errors == ()
    assert [
        (
            row.source_sku,
            row.unit_price_ex_vat,
            row.effective_from,
            row.effective_to,
        )
        for row in parsed.candidates
    ] == [
        (
            "00123",
            Decimal("100.250000000000"),
            date(2025, 1, 1),
            date(2025, 1, 31),
        ),
        ("123", Decimal("200.000000000000"), date(2025, 2, 1), None),
    ]


def test_parse_dh_price_workbook_reports_all_blocking_row_errors() -> None:
    parsed = parse_dh_price_workbook(
        _workbook_bytes(
            [
                ("A", 10, date(2025, 1, 1), None),
                ("A", 10, date(2025, 1, 1), None),
                ("B", 0, date(2025, 1, 1), None),
                ("C", 20, date(2025, 2, 1), date(2025, 1, 31)),
                ("D", 30, "01/02/2025", None),
                ("E", "10000000000000000", date(2025, 1, 1), None),
            ]
        )
    )

    assert parsed.candidates == ()
    assert any("SKU A" in error and "ซ้ำ" in error for error in parsed.errors)
    assert any("แถว 4" in error and "มากกว่า 0" in error for error in parsed.errors)
    assert any("แถว 5" in error and "ช่วงวันที่" in error for error in parsed.errors)
    assert any("แถว 6" in error and "YYYY-MM-DD" in error for error in parsed.errors)
    assert any("แถว 7" in error and "เกิน" in error for error in parsed.errors)


def test_parse_dh_price_workbook_rejects_extra_headers_and_empty_data() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "DH Price Master"
    sheet.append((*HEADERS, "Unexpected"))
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(DhPriceMasterError, match="Column ต้องเรียง"):
        parse_dh_price_workbook(output.getvalue())

    empty = parse_dh_price_workbook(_workbook_bytes([]))
    assert empty.candidates == ()
    assert any("ไม่พบข้อมูลราคา" in error for error in empty.errors)


def test_preview_is_incremental_and_does_not_mutate_existing_prices() -> None:
    engine, session = _database()
    try:
        session.add(ModernTrade(id=1, code="DH", name="DoHome", vat_mode="exclude"))
        session.add_all(
            [
                _price(1, 1, "A", 100, date(2025, 1, 1), date(2025, 1, 31)),
                _price(2, 1, "B", 200, date(2025, 1, 1)),
                _price(3, 1, "KEEP", 300, date(2025, 1, 1)),
            ]
        )
        session.commit()

        preview = preview_dh_price_master(
            session,
            _workbook_bytes(
                [
                    ("A", 110, date(2025, 1, 1), date(2025, 1, 31)),
                    ("A", 120, date(2025, 2, 1), None),
                    ("B", 200, date(2025, 1, 1), None),
                ]
            ),
        )

        assert preview.inserted == 1
        assert preview.updated == 1
        assert preview.unchanged == 1
        assert preview.errors == ()
        assert len(session.scalars(select(DhEffectivePrice)).all()) == 3
        assert session.get(DhEffectivePrice, 1).unit_price_ex_vat == 100
        assert session.scalar(
            select(DhEffectivePrice).where(DhEffectivePrice.source_sku == "KEEP")
        )
    finally:
        session.close()
        engine.dispose()


def test_preview_rejects_overlap_with_retained_database_row() -> None:
    engine, session = _database()
    try:
        session.add(ModernTrade(id=1, code="DH", name="DoHome", vat_mode="exclude"))
        session.add(_price(1, 1, "A", 100, date(2025, 1, 1)))
        session.commit()

        preview = preview_dh_price_master(
            session,
            _workbook_bytes([("A", 120, date(2025, 2, 1), None)]),
        )

        assert preview.inserted == 1
        assert any("SKU A" in error and "ซ้อน" in error for error in preview.errors)
    finally:
        session.close()
        engine.dispose()


def test_load_effective_prices_respects_boundaries_and_dh_scope() -> None:
    engine, session = _database()
    try:
        session.add_all(
            [
                ModernTrade(id=1, code="DH", name="DoHome", vat_mode="exclude"),
                ModernTrade(id=2, code="GH", name="Global House"),
            ]
        )
        session.add_all(
            [
                _price(1, 1, "A", 100, date(2025, 1, 1), date(2025, 1, 31)),
                _price(2, 1, "A", 110, date(2025, 2, 1)),
                _price(3, 2, "A", 999, date(2025, 1, 1)),
            ]
        )
        session.commit()

        january = load_effective_dh_prices(session, date(2025, 1, 31), {"A"})
        february = load_effective_dh_prices(session, date(2025, 2, 1), {"A"})

        assert [(row.source_sku, row.unit_price_ex_vat) for row in january] == [
            ("A", Decimal("100.000000000000"))
        ]
        assert [(row.source_sku, row.unit_price_ex_vat) for row in february] == [
            ("A", Decimal("110.000000000000"))
        ]
    finally:
        session.close()
        engine.dispose()


def test_dh_price_model_has_incremental_key_and_database_checks() -> None:
    table = inspect(DhEffectivePrice).local_table
    unique_columns = {
        tuple(column.name for column in constraint.columns)
        for constraint in table.constraints
        if constraint.__class__.__name__ == "UniqueConstraint"
    }
    check_sql = {
        str(constraint.sqltext)
        for constraint in table.constraints
        if hasattr(constraint, "sqltext")
    }

    assert ("modern_trade_id", "source_sku", "effective_from") in unique_columns
    assert "unit_price_ex_vat > 0" in check_sql
    assert "effective_to IS NULL OR effective_to >= effective_from" in check_sql
