from datetime import UTC, date, datetime
from decimal import Decimal
from itertools import count
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.api.performance import performance
from app.database import Base
from app.importers.gh import GhFormatError, extract_gh_file, inspect_gh_workbook
from app.models import (
    BranchMapping,
    InventoryCoverage,
    ItemMapping,
    ModernTrade,
    SalesInventoryFact,
    SourceFile,
)
from app.services import gh_import
from app.services.automatic_import import SourceCandidate
from app.services.gh_automatic_import import _gh_unchanged_outcome
from app.services.gh_import import GhImportError, import_gh_file

HEADERS = [
    "Product Code",
    "Product Name",
    "Product Status",
    "Branch Code",
    "Branch Name",
    "Stock (Qty)",
    "Stock Amount (Ex VAT)",
    "Stock Amount (In VAT)",
    "Cost (Ex VAT)",
    "Cost (In VAT)",
    "Sale Quantity",
    "Sale Amount",
]


def _report(path: Path, *, footer_sale_amount: object = 0) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    sheet.append(
        [
            "00001234",
            "Window A",
            "A",
            "GH-101",
            "สำนักงานใหญ่",
            2,
            200,
            214,
            100,
            107,
            1,
            107,
        ]
    )
    sheet.append(
        [
            "SKU-A7",
            "Door B",
            "I",
            "GH-102",
            "สาขาขอนแก่น",
            0,
            0,
            0,
            0,
            0,
            -1,
            -107,
        ]
    )
    sheet.append([None, None, None, None, None, None, 200, 214, 100, 107, None, footer_sale_amount])
    workbook.save(path)


def _legacy_report(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"

    sheet.merge_cells("H2:L2")
    sheet["H2"] = "GH-003"
    sheet.merge_cells("M2:O2")
    sheet["M2"] = "GH-101"
    sheet.merge_cells("P2:R3")
    sheet["P2"] = "Grand Total"
    sheet.merge_cells("H3:L3")
    sheet["H3"] = "DC2"
    sheet.merge_cells("M3:O3")
    sheet["M3"] = "RE"

    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Product Code"
    sheet.merge_cells("C4:E4")
    sheet["C4"] = "Product Name"
    sheet["F4"] = "Unit"
    sheet["G4"] = "Inactive"
    sheet.merge_cells("H4:I4")
    sheet["H4"] = "Stock"
    sheet.merge_cells("J4:K4")
    sheet["J4"] = "Sale Quantity"
    sheet["L4"] = "Net Sale"
    for start in (13, 16):
        sheet.cell(4, start, "Stock")
        sheet.cell(4, start + 1, "Sale Quantity")
        sheet.cell(4, start + 2, "Net Sale")

    rows = [
        (5, "00001234", "Window A", "A", (2, 1, 107), (3, 0, 0), (5, 1, 107)),
        (6, "SKU-A7", "Door B", "I", (0, 0, 0), (0, -1, -107), (0, -1, -107)),
    ]
    for row_number, sku, description, status, first, second, total in rows:
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=2)
        sheet.cell(row_number, 1, sku)
        sheet.merge_cells(start_row=row_number, start_column=3, end_row=row_number, end_column=5)
        sheet.cell(row_number, 3, description)
        sheet.cell(row_number, 7, status)
        sheet.merge_cells(start_row=row_number, start_column=8, end_row=row_number, end_column=9)
        sheet.cell(row_number, 8, first[0])
        sheet.merge_cells(start_row=row_number, start_column=10, end_row=row_number, end_column=11)
        sheet.cell(row_number, 10, first[1])
        sheet.cell(row_number, 12, first[2])
        for offset, value in enumerate(second):
            sheet.cell(row_number, 13 + offset, value)
        for offset, value in enumerate(total):
            sheet.cell(row_number, 16 + offset, value)

    sheet.merge_cells("A7:G7")
    sheet["A7"] = "Grand Total"
    sheet.merge_cells("H7:I7")
    sheet["H7"] = 2
    sheet.merge_cells("J7:K7")
    sheet["J7"] = 1
    sheet["L7"] = 107
    for column, value in enumerate((3, -1, -107, 5, 0, 0), start=13):
        sheet.cell(7, column, value)
    workbook.save(path)


def _empty_legacy_report(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["H2"] = "Grand Total"
    sheet["G3"] = "Inactive"
    sheet["A4"] = "Grand Total"
    workbook.save(path)


def test_extract_gh_file_uses_filename_date_and_preserves_source_metrics(tmp_path: Path) -> None:
    source = tmp_path / "Piyawat-2026-09-12-074128.xlsx"
    _report(source)

    extract = extract_gh_file(source)

    assert extract.data_date == date(2026, 9, 11)
    assert extract.sales_grain == "rolling_30d"
    assert extract.sales_window_days == 30
    assert extract.summary.row_count == 2
    assert extract.summary.store_count == 2
    assert extract.summary.sku_count == 2
    assert extract.summary.negative_row_count == 1
    assert extract.summary.source_amount == Decimal("0")
    assert extract.summary.amount == Decimal("0E-12")
    assert extract.summary.sales_qty == Decimal("0")
    assert extract.summary.stock_on_hand == Decimal("2")
    assert extract.summary.stock_value == Decimal("200")
    assert extract.inventory_skus == frozenset({"00001234", "SKU-A7"})
    assert extract.inventory_branches == frozenset({"GH-101", "GH-102"})
    assert extract.product_status_counts == (("A", 1), ("I", 1))
    assert extract.rows[0].sku == "00001234"
    assert extract.rows[0].unit_cost_ex_vat == Decimal("100")
    assert extract.rows[0].stock_amount_in_vat == Decimal("214")
    assert extract.rows[0].product_status == "A"
    assert extract.reconciliation_errors == ()


def test_extract_gh_file_supports_legacy_wide_branch_layout(tmp_path: Path) -> None:
    source = tmp_path / "Piyawat-2026-08-03-080548.xlsx"
    _legacy_report(source)

    extract = extract_gh_file(source)

    assert extract.data_date == date(2026, 8, 2)
    assert extract.sales_grain == "daily"
    assert extract.sales_window_days is None
    assert extract.summary.row_count == 3
    assert extract.summary.store_count == 2
    assert extract.summary.sku_count == 2
    assert extract.summary.negative_row_count == 1
    assert extract.summary.source_amount == Decimal("0")
    assert extract.summary.amount == Decimal("0E-12")
    assert extract.summary.sales_qty == Decimal("0")
    assert extract.summary.stock_on_hand == Decimal("5")
    assert extract.summary.stock_value == Decimal("0")
    assert extract.inventory_skus == frozenset({"00001234", "SKU-A7"})
    assert extract.inventory_branches == frozenset({"GH-003", "GH-101"})
    assert extract.product_status_counts == (("A", 2), ("I", 1))
    assert extract.footer_totals == (
        ("Stock (Qty)", Decimal("5")),
        ("Sale Quantity", Decimal("0")),
        ("Sale Amount", Decimal("0")),
    )
    assert extract.reconciliation_errors == ()


def test_extract_gh_file_supports_empty_legacy_report(tmp_path: Path) -> None:
    source = tmp_path / "Piyawat-2025-04-03-082519.xlsx"
    _empty_legacy_report(source)

    extract = extract_gh_file(source)

    assert extract.data_date == date(2025, 4, 2)
    assert extract.rows == ()
    assert extract.summary.row_count == 0
    assert extract.summary.store_count == 0
    assert extract.summary.sku_count == 0
    assert extract.summary.source_amount == Decimal("0")
    assert extract.summary.sales_qty == Decimal("0")
    assert extract.summary.stock_on_hand == Decimal("0")
    assert extract.inventory_skus == frozenset()
    assert extract.inventory_branches == frozenset()
    assert extract.reconciliation_errors == ()


def test_inspect_gh_workbook_requires_filename_and_signature(tmp_path: Path) -> None:
    source = tmp_path / "Piyawat-2026-01-01.xlsx"
    _report(source)

    assert inspect_gh_workbook(source) == ("combined", date(2025, 12, 31))

    renamed = tmp_path / "Other-2026-01-01.xlsx"
    source.rename(renamed)
    with pytest.raises(GhFormatError, match="Piyawat"):
        inspect_gh_workbook(renamed)


def test_extract_gh_file_rejects_duplicate_branch_sku_grain(tmp_path: Path) -> None:
    source = tmp_path / "Piyawat-2026-09-12.xlsx"
    _report(source)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    row = ["00001", "Window", "A", "GH-101", "สำนักงานใหญ่", 1, 100, 107, 100, 107, 0, 0]
    sheet.append(row)
    sheet.append(row)
    sheet.append([None, None, None, None, None, None, 200, 214, 200, 214, None, 0])
    workbook.save(source)

    with pytest.raises(GhFormatError, match="SKU × Branch ซ้ำ"):
        extract_gh_file(source)


def test_extract_gh_file_reports_footer_mismatch(tmp_path: Path) -> None:
    source = tmp_path / "Piyawat-2026-09-12.xlsx"
    _report(source, footer_sale_amount=999)

    extract = extract_gh_file(source)

    assert extract.summary.source_amount == Decimal("0")
    assert any("Sale Amount" in warning for warning in extract.reconciliation_errors)


def test_extract_gh_file_rejects_wrong_headers(tmp_path: Path) -> None:
    source = tmp_path / "Piyawat-2026-09-12.xlsx"
    _report(source)
    workbook = Workbook()
    workbook.active.append(["Product Code", "Wrong Header"])
    workbook.save(source)

    with pytest.raises(GhFormatError, match="โครงสร้าง column"):
        extract_gh_file(source)


def test_import_gh_file_isolated_trade_and_source_audit(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "Piyawat-2026-09-12.xlsx"
    _report(source)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    batch_ids = count(1)
    fact_ids = count(1)
    coverage_ids = count(1)
    monkeypatch.setattr(
        gh_import,
        "_new_batch",
        lambda **values: gh_import.ImportBatch(id=next(batch_ids), **values),
    )
    monkeypatch.setattr(
        gh_import,
        "_new_fact",
        lambda **values: gh_import.SalesInventoryFact(id=next(fact_ids), **values),
    )
    monkeypatch.setattr(
        gh_import,
        "_new_coverage",
        lambda **values: gh_import.InventoryCoverage(id=next(coverage_ids), **values),
    )

    with Session(engine) as session:
        other = ModernTrade(code="TWD", name="Thai Watsadu", vat_mode="include")
        session.add(other)
        session.flush()
        batch = import_gh_file(session, extract_gh_file(source), actor="admin")
        session.commit()

        gh = session.scalar(select(ModernTrade).where(ModernTrade.code == "GH"))
        assert gh is not None
        assert gh.name == "Global House"
        assert gh.source_subfolder == "GBH"
        assert gh.source_enabled is False
        assert gh.schedule_enabled is False
        assert gh.show_unmatched_items is False
        assert gh.show_unmatched_branches is False
        assert batch.status == "imported"
        assert '"productStatusCounts": {"A": 1, "I": 1}' in (batch.source_pair_json or "")
        assert (
            session.scalar(
                select(func.count(SalesInventoryFact.id)).where(
                    SalesInventoryFact.modern_trade_id == gh.id
                )
            )
            == 2
        )
        assert (
            session.scalar(
                select(func.count(SalesInventoryFact.id)).where(
                    SalesInventoryFact.modern_trade_id == other.id
                )
            )
            == 0
        )
        coverage = session.scalar(
            select(InventoryCoverage).where(InventoryCoverage.modern_trade_id == gh.id)
        )
        assert coverage is not None
        assert coverage.source_row_count == 2


def test_import_gh_file_records_empty_legacy_day(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "Piyawat-2025-04-03-082519.xlsx"
    _empty_legacy_report(source)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(
        gh_import,
        "_new_batch",
        lambda **values: gh_import.ImportBatch(id=1, **values),
    )
    monkeypatch.setattr(
        gh_import,
        "_new_coverage",
        lambda **values: gh_import.InventoryCoverage(id=1, **values),
    )

    with Session(engine) as session:
        batch = import_gh_file(session, extract_gh_file(source), actor="admin")
        session.commit()

        gh = session.scalar(select(ModernTrade).where(ModernTrade.code == "GH"))
        assert gh is not None
        assert batch.status == "imported"
        assert batch.row_count == 0
        assert (
            session.scalar(
                select(func.count(SalesInventoryFact.id)).where(
                    SalesInventoryFact.modern_trade_id == gh.id
                )
            )
            == 0
        )
        coverage = session.scalar(
            select(InventoryCoverage).where(InventoryCoverage.modern_trade_id == gh.id)
        )
        assert coverage is not None
        assert coverage.source_row_count == 0


def test_gh_automatic_import_retries_unchanged_failed_source() -> None:
    modified_at = datetime(2026, 9, 14, tzinfo=UTC)
    source = SourceFile(
        modern_trade_id=1,
        source_path=r"\\server\GBH\Piyawat-2025-01-01.xlsx",
        source_filename="Piyawat-2025-01-01.xlsx",
        size_bytes=100,
        modified_at=modified_at,
        status="failed",
        error_message="legacy layout was unsupported",
    )
    candidate = SourceCandidate(
        path=source.source_path,
        filename=source.source_filename,
        size_bytes=source.size_bytes,
        modified_at=modified_at,
    )

    assert _gh_unchanged_outcome(source, candidate) is None


def test_import_gh_file_rejects_same_business_fingerprint(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "Piyawat-2026-09-12.xlsx"
    _report(source)
    extract = extract_gh_file(source)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    batch_ids = count(1)
    fact_ids = count(1)
    coverage_ids = count(1)
    monkeypatch.setattr(
        gh_import,
        "_new_batch",
        lambda **values: gh_import.ImportBatch(id=next(batch_ids), **values),
    )
    monkeypatch.setattr(
        gh_import,
        "_new_fact",
        lambda **values: gh_import.SalesInventoryFact(id=next(fact_ids), **values),
    )
    monkeypatch.setattr(
        gh_import,
        "_new_coverage",
        lambda **values: gh_import.InventoryCoverage(id=next(coverage_ids), **values),
    )

    with Session(engine) as session:
        import_gh_file(session, extract)
        session.commit()
        with pytest.raises(GhImportError, match="ถูกนำเข้าแล้ว"):
            import_gh_file(session, extract)


def test_local_compose_applies_migrations_before_starting_api() -> None:
    compose = (Path(__file__).resolve().parents[2] / "compose.yaml").read_text(
        encoding="utf-8"
    )

    assert "python -m alembic upgrade head && exec uvicorn" in compose


def test_gh_latest_branch_month_with_mappings_and_no_facts_returns_empty_state() -> None:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add_all(
            [
                ModernTrade(id=1, code="GH", name="Global House"),
                ItemMapping(
                    id=1,
                    modern_trade_id=1,
                    source_sku="8859283002063",
                    source_description="GH item",
                    wa_item_code="WA-GH",
                    status="confirmed",
                    effective_from=date(2025, 1, 1),
                    changed_by="test",
                ),
                BranchMapping(
                    id=1,
                    modern_trade_id=1,
                    source_branch_code="GH-101",
                    source_branch_description="GH branch",
                    wa_branch_code="CGH-0001",
                    wa_branch_description="Global House 101",
                    status="confirmed",
                    effective_from=date(2025, 1, 1),
                    changed_by="test",
                ),
            ]
        )
        session.commit()

        report = performance(
            session=session,
            mt_code="GH",
            date_from=None,
            date_to=None,
            date_range=None,
            page=1,
            page_size=25,
            branch_id=None,
            branch_ids=None,
            sku_ids=None,
            sku_flag="all",
            mapping_status=None,
            hide_unmapped=False,
            search=None,
            grain="branch_month",
            period_month="latest",
            latest_only=False,
            sales_basis="net",
            include_turnover=False,
            report_mode="sales",
        )

    assert report["selectedMonth"] is None
    assert report["summary"]["amount"] == 0
    assert report["meta"]["totalSkus"] == 1
    assert report["branches"] == [
        {"id": "GH-101", "name": "Global House 101"},
    ]
