from datetime import date
from decimal import Decimal
from itertools import count
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.database import Base
from app.importers.hh import HhFormatError, extract_hh_pair, inspect_hh_workbook
from app.models import ImportBatch, ModernTrade, SalesInventoryFact
from app.services import hh_import


def test_fact_taxonomy_columns_support_homehub_labels() -> None:
    assert SalesInventoryFact.__table__.c.category.type.length == 300
    assert SalesInventoryFact.__table__.c.subcategory.type.length == 300


def _report(path: Path, *, kind: str, report_date: str = "10-09-2026") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, f"Report date {report_date} generated 2026-09-11 07:30:00")
    headings = ["#", "Category", "Group", "SKU", "Description", "Unit"]
    branches = ["Ubon", "Chayangkun", "Warin", "Khon Kaen", "Amnat", "Total"]
    for column, heading in enumerate(headings, 1):
        sheet.cell(2, column, heading)
    for index, branch in enumerate(branches):
        sheet.cell(2, 7 + index * 2, branch)
        sheet.cell(3, 7 + index * 2, "Qty")
        sheet.cell(3, 8 + index * 2, "Value")
    rows = [
        ("00001", "Leading zero item"),
        ("12345678901", "Eleven digit item"),
        ("SKU-A7", "Opaque item"),
    ]
    for row_number, (sku, description) in enumerate(rows, 4):
        sheet.cell(row_number, 1, row_number - 3)
        sheet.cell(row_number, 2, "Category")
        sheet.cell(row_number, 3, "Group")
        sheet.cell(row_number, 4, sku)
        sheet.cell(row_number, 5, description)
        for branch_index in range(5):
            qty_column = 7 + branch_index * 2
            sheet.cell(row_number, qty_column, row_number + branch_index)
            sheet.cell(row_number, qty_column + 1, (row_number + branch_index) * 10)
    workbook.save(path)


def test_extract_hh_pair_preserves_opaque_skus_and_all_branches(tmp_path: Path) -> None:
    stock = tmp_path / "StockReport.xlsx"
    sales = tmp_path / "SaleReport.xlsx"
    _report(stock, kind="stock")
    _report(sales, kind="sales")

    extract = extract_hh_pair(stock, sales)

    assert extract.data_date.isoformat() == "2026-09-10"
    assert extract.summary.sku_count == 3
    assert extract.summary.store_count == 5
    assert extract.summary.row_count == 15
    assert {row.sku for row in extract.rows} == {"00001", "12345678901", "SKU-A7"}
    assert extract.summary.amount == Decimal("1050.000000000000")
    assert extract.summary.sales_qty == Decimal("105")
    assert extract.summary.stock_on_hand == Decimal("105")
    assert extract.summary.stock_value == Decimal("1050")


def test_extract_hh_pair_rejects_different_report_dates(tmp_path: Path) -> None:
    stock = tmp_path / "StockReport.xlsx"
    sales = tmp_path / "SaleReport.xlsx"
    _report(stock, kind="stock")
    _report(sales, kind="sales", report_date="11-09-2026")

    with pytest.raises(HhFormatError, match="วันที่ข้อมูล Stock และ Sale ไม่ตรงกัน"):
        extract_hh_pair(stock, sales)


def test_inspect_hh_workbook_uses_structure_and_report_kind(tmp_path: Path) -> None:
    stock = tmp_path / "2026-09-10-StockReport.xlsx"
    sales = tmp_path / "2026-09-10-SaleReport.xlsx"
    _report(stock, kind="stock")
    _report(sales, kind="sales")

    assert inspect_hh_workbook(stock) == ("inventory", date(2026, 9, 10))
    assert inspect_hh_workbook(sales) == ("sales", date(2026, 9, 10))


def test_hh_backfill_adapter_adds_only_the_requested_opaque_sku(
    tmp_path: Path, monkeypatch
) -> None:
    stock = tmp_path / "StockReport.xlsx"
    sales = tmp_path / "SaleReport.xlsx"
    _report(stock, kind="stock")
    _report(sales, kind="sales")
    pair = extract_hh_pair(stock, sales)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    fact_ids = count(1)
    fact_factory = hh_import._new_fact
    monkeypatch.setattr(
        hh_import,
        "_new_fact",
        lambda **values: fact_factory(id=next(fact_ids), **values),
    )

    with Session(engine) as session:
        session.add(
            ModernTrade(
                id=1,
                code="HH",
                name="HomeHub",
                show_unmatched_items=True,
                show_unmatched_branches=True,
            )
        )
        session.add(
            ImportBatch(
                id=1,
                modern_trade_id=1,
                status="imported",
                data_date=pair.data_date,
                source_path="manual:StockReport.xlsx",
                source_filename="StockReport.xlsx + SaleReport.xlsx",
                checksum_sha256="a" * 64,
                row_count=0,
                store_count=5,
                sku_count=3,
                negative_row_count=0,
                source_amount=0,
                amount=0,
                sales_qty=0,
                stock_on_hand=0,
                reported_stock_on_hand=0,
                stock_on_order=0,
                stock_value=0,
            )
        )
        session.commit()
        batch = session.get(ImportBatch, 1)
        assert batch is not None

        inserted = hh_import.append_hh_sku_facts(session, batch, pair, "SKU-A7")
        session.commit()
        stored_skus = set(session.scalars(select(SalesInventoryFact.source_sku)))
        fact_count = session.scalar(select(func.count()).select_from(SalesInventoryFact))

    assert inserted == 5
    assert fact_count == 5
    assert stored_skus == {"SKU-A7"}
