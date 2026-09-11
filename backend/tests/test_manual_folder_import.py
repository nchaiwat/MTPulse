from datetime import UTC, datetime
from decimal import Decimal
from itertools import count
from types import SimpleNamespace

from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app import worker
from app.database import Base
from app.importers.twd import TwdExtract, TwdSummary
from app.models import AuditEvent, ManualUploadBatch, ManualUploadFile
from app.services import manual_upload_batches
from app.services.manual_upload_batches import (
    create_folder_batch,
    finalize_folder_batch,
    process_folder_batch,
    queue_folder_batch,
    store_folder_file,
)


def _settings(path):
    return SimpleNamespace(manual_upload_staging_dir=str(path))


def _detection(group: str = "TWD"):
    return SimpleNamespace(
        status="detected",
        source_group_code=group,
        mt_codes=(group,) if group == "TWD" else ("HP", "MH"),
        source_kind="workbook" if group == "TWD" else "inventory",
        reason=None,
    )


def _hh_report(path, *, report_date: str = "10-09-2026") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, f"Report date {report_date}")
    for column, heading in enumerate(["#", "Category", "Group", "SKU", "Description", "Unit"], 1):
        sheet.cell(2, column, heading)
    for index, branch in enumerate(["Ubon", "Chayangkun", "Warin", "Khon Kaen", "Amnat", "Total"]):
        sheet.cell(2, 7 + index * 2, branch)
        sheet.cell(3, 7 + index * 2, "Qty")
        sheet.cell(3, 8 + index * 2, "Value")
    for row_number, sku in enumerate(["00001", "SKU-A7"], 4):
        sheet.cell(row_number, 1, row_number - 3)
        sheet.cell(row_number, 2, "Category")
        sheet.cell(row_number, 3, "Group")
        sheet.cell(row_number, 4, sku)
        sheet.cell(row_number, 5, f"Product {sku}")
        for branch_index in range(5):
            sheet.cell(row_number, 7 + branch_index * 2, row_number + branch_index)
            sheet.cell(row_number, 8 + branch_index * 2, (row_number + branch_index) * 10)
    workbook.save(path)


def test_folder_batch_uploads_independently_and_finalizes_one_mt(tmp_path, monkeypatch) -> None:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(manual_upload_batches, "get_settings", lambda: _settings(tmp_path))
    monkeypatch.setattr(manual_upload_batches, "detect_upload_file", lambda _path: _detection())

    with Session(engine) as session:
        batch = create_folder_batch(session, file_count=2, actor="admin")
        first = store_folder_file(
            session,
            batch,
            filename="one.xls",
            content=b"one",
            idempotency_key="one",
        )
        repeated = store_folder_file(
            session,
            batch,
            filename="one.xls",
            content=b"one",
            idempotency_key="one",
        )
        store_folder_file(
            session,
            batch,
            filename="two.xls",
            content=b"two",
            idempotency_key="two",
        )
        finalized = finalize_folder_batch(session, batch, expected_source_group="TWD")
        first_id = first.id
        repeated_id = repeated.id

    assert first_id == repeated_id
    assert finalized.status == "awaiting_confirmation"
    assert finalized.detected_source_group == "TWD"
    assert finalized.uploaded_count == 2
    assert finalized.eligible_count == 2


def test_folder_batch_blocks_wrong_mt_before_confirmation(tmp_path, monkeypatch) -> None:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(manual_upload_batches, "get_settings", lambda: _settings(tmp_path))
    monkeypatch.setattr(
        manual_upload_batches,
        "detect_upload_file",
        lambda _path: _detection("HP_MH"),
    )

    with Session(engine) as session:
        batch = create_folder_batch(session, file_count=1, actor="admin")
        store_folder_file(
            session,
            batch,
            filename="inventory.zip",
            content=b"zip",
            idempotency_key="one",
        )
        finalized = finalize_folder_batch(session, batch, expected_source_group="TWD")

    assert finalized.status == "failed"
    assert finalized.detection_status == "conflict"
    assert "เลือกหน้า TWD" in (finalized.error_message or "")


def test_folder_worker_continues_after_one_invalid_twd_file(tmp_path, monkeypatch) -> None:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(manual_upload_batches, "get_settings", lambda: _settings(tmp_path))
    monkeypatch.setattr(manual_upload_batches, "detect_upload_file", lambda _path: _detection())
    ids = count(1)
    monkeypatch.setattr(
        manual_upload_batches,
        "AuditEvent",
        lambda **values: AuditEvent(id=next(ids), **values),
    )

    def extract(path):
        if "bad" in path.name:
            raise ValueError("invalid workbook")
        summary = TwdSummary(
            row_count=0,
            store_count=0,
            sku_count=0,
            negative_row_count=0,
            source_amount=Decimal("0"),
            amount=Decimal("0"),
            sales_qty=Decimal("0"),
            stock_on_hand=Decimal("0"),
            stock_on_order=Decimal("0"),
        )
        return TwdExtract(
            source_path=str(path),
            source_filename=path.name,
            checksum_sha256="a" * 64,
            data_date=datetime(2026, 9, 9, tzinfo=UTC).date(),
            rows=(),
            summary=summary,
            reported_summary=summary,
            reconciliation_errors=(),
        )

    monkeypatch.setattr(manual_upload_batches, "extract_twd_file", extract)
    monkeypatch.setattr(
        manual_upload_batches,
        "import_twd_extract",
        lambda _session, _extract: SimpleNamespace(id=99),
    )

    with Session(engine) as session:
        batch = create_folder_batch(session, file_count=2, actor="admin")
        store_folder_file(
            session,
            batch,
            filename="good.xls",
            content=b"good",
            idempotency_key="good",
        )
        store_folder_file(
            session,
            batch,
            filename="bad.xls",
            content=b"bad",
            idempotency_key="bad",
        )
        finalize_folder_batch(session, batch, expected_source_group="TWD")
        queue_folder_batch(session, batch)
        process_folder_batch(session, batch.id)
        stored_batch = session.get(ManualUploadBatch, batch.id)
        rows = session.query(ManualUploadFile).order_by(ManualUploadFile.id).all()

    assert stored_batch is not None
    assert stored_batch.status == "completed_with_issues"
    assert [row.status for row in rows] == ["imported", "failed"]


def test_worker_requeues_interrupted_folder_batch(tmp_path, monkeypatch) -> None:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    monkeypatch.setattr(manual_upload_batches, "get_settings", lambda: _settings(tmp_path))
    monkeypatch.setattr(worker, "SessionLocal", factory)

    with factory() as session:
        batch = create_folder_batch(session, file_count=1, actor="admin")
        batch.status = "processing"
        session.commit()
        batch_id = batch.id

    assert worker.recover_interrupted_manual_upload_batches() == 1
    with factory() as session:
        stored = session.get(ManualUploadBatch, batch_id)
        assert stored is not None
        assert stored.status == "queued"


def test_hp_mh_folder_completion_is_visible_in_both_mt_audit_streams(
    tmp_path, monkeypatch
) -> None:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(manual_upload_batches, "get_settings", lambda: _settings(tmp_path))
    ids = count(1)
    monkeypatch.setattr(
        manual_upload_batches,
        "AuditEvent",
        lambda **values: AuditEvent(id=next(ids), **values),
    )

    with Session(engine) as session:
        batch = create_folder_batch(session, file_count=1, actor="admin")
        batch.detected_source_group = "HP_MH"
        manual_upload_batches._finish_batch(session, batch)
        events = session.query(AuditEvent).order_by(AuditEvent.id).all()

    payloads = [event.after_json or "" for event in events]
    assert any('"mtCode": "HP"' in payload for payload in payloads)
    assert any('"mtCode": "MH"' in payload for payload in payloads)


def test_hh_folder_import_pairs_stock_and_sales_without_sku_pattern(
    tmp_path, monkeypatch
) -> None:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(manual_upload_batches, "get_settings", lambda: _settings(tmp_path))
    stock = tmp_path / "source-StockReport.xlsx"
    sales = tmp_path / "source-SaleReport.xlsx"
    _hh_report(stock)
    _hh_report(sales)
    imported_pairs = []
    audit_ids = count(1)
    monkeypatch.setattr(
        manual_upload_batches,
        "AuditEvent",
        lambda **values: AuditEvent(id=next(audit_ids), **values),
    )

    def import_pair(_session, pair, *, actor):
        imported_pairs.append((pair, actor))
        return SimpleNamespace(id=77)

    monkeypatch.setattr(manual_upload_batches, "import_hh_pair", import_pair)

    with Session(engine) as session:
        batch = create_folder_batch(session, file_count=2, actor="admin")
        store_folder_file(
            session,
            batch,
            filename="2026-09-10-StockReport.xlsx",
            content=stock.read_bytes(),
            idempotency_key="hh-stock",
        )
        store_folder_file(
            session,
            batch,
            filename="2026-09-10-SaleReport.xlsx",
            content=sales.read_bytes(),
            idempotency_key="hh-sales",
        )
        finalize_folder_batch(session, batch, expected_source_group="HH")
        queue_folder_batch(session, batch)
        process_folder_batch(session, batch.id)
        stored_batch = session.get(ManualUploadBatch, batch.id)
        files = session.query(ManualUploadFile).order_by(ManualUploadFile.id).all()

    assert stored_batch is not None
    assert stored_batch.status == "completed"
    assert stored_batch.imported_count == 2
    assert [row.status for row in files] == ["imported", "imported"]
    assert [row.source_kind for row in files] == ["inventory", "sales"]
    assert {row.sku for row in imported_pairs[0][0].rows} == {"00001", "SKU-A7"}
    assert imported_pairs[0][1] == "admin"
