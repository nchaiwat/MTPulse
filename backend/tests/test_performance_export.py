from datetime import UTC, datetime
from io import BytesIO

import openpyxl
import pytest

from app.services.performance_export import (
    build_performance_workbook,
    performance_export_filename,
)


def _report() -> dict:
    return {
        "dates": ["2026-08-17"],
        "selectedMonth": "2026-08",
        "branches": [
            {"id": "60016", "name": "ภูเก็ต เฟสติวัล"},
            {"id": "60926", "name": "ลำปาง"},
        ],
        "items": [
            {
                "sku": "060277546",
                "isSho": True,
                "isPro": False,
                "twdDescription": "สินค้าไทยหนึ่ง",
                "waItem": "FA09-W0112-080050",
                "waDescription": "สินค้า WA หนึ่ง",
                "points": [
                    {
                        "date": "2026-08-17",
                        "branchId": "60016",
                        "amount": 100.5,
                        "qty": 2,
                        "stockOh": 7,
                        "stockOnOrder": 1,
                    },
                    {
                        "date": "2026-08-17",
                        "branchId": "60926",
                        "amount": 50,
                        "qty": 1,
                        "stockOh": 3,
                        "stockOnOrder": 0,
                    },
                ],
            },
            {
                "sku": "060277547",
                "isSho": False,
                "isPro": True,
                "twdDescription": "สินค้าไทยสอง",
                "waItem": "FA09-W0112-100110",
                "waDescription": "สินค้า WA สอง",
                "points": [
                    {
                        "date": "2026-08-17",
                        "branchId": "60926",
                        "amount": -25,
                        "qty": -1,
                        "stockOh": 4,
                        "stockOnOrder": 2,
                    },
                ],
            },
        ],
    }


def test_performance_workbook_contains_all_rows_formulas_and_branch_headers() -> None:
    content = build_performance_workbook(
        _report(),
        mode="sales",
        metric="amount",
        grain="day",
        show_descriptions=True,
        branch_id=None,
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Report"]

    assert sheet["A5"].value == "Sho"
    assert sheet["B5"].value == "Pro"
    assert sheet["C5"].value == "TWD SKU"
    assert sheet["D5"].value == "TWD Description"
    assert sheet["G5"].value == "TOTAL"
    assert sheet["H5"].value == "60016 - ภูเก็ต เฟสติวัล"
    assert sheet["I5"].value == "60926 - ลำปาง"
    assert sheet["A6"].value == "✓"
    assert sheet["B7"].value == "✓"
    assert sheet["C6"].value == "060277546"
    assert sheet["C7"].value == "060277547"
    assert sheet["G6"].value == "=SUM(H6:I6)"
    assert sheet["G4"].value == "=SUM(G6:G7)"
    assert sheet["H6"].value == 100.5
    assert sheet["I6"].value == 50
    assert sheet["I7"].value == -25
    assert sheet.freeze_panes == "H6"
    assert sheet.auto_filter.ref == "A5:I7"
    workbook.close()


def test_performance_workbook_can_hide_description_columns() -> None:
    content = build_performance_workbook(
        _report(),
        mode="sales",
        metric="qty",
        grain="day",
        show_descriptions=False,
        branch_id=None,
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Report"]

    assert [sheet.cell(5, column).value for column in range(1, 8)] == [
        "Sho",
        "Pro",
        "TWD SKU",
        "WA Item",
        "TOTAL",
        "60016 - ภูเก็ต เฟสติวัล",
        "60926 - ลำปาง",
    ]
    assert sheet["E6"].value == "=SUM(F6:G6)"
    workbook.close()




def test_performance_workbook_keeps_backend_aggregate_for_selected_branches() -> None:
    report = _report()
    report["items"] = [
        {
            **report["items"][0],
            "points": [
                {
                    "date": "2026-08-17",
                    "branchId": "all",
                    "amount": 150.5,
                    "qty": 3,
                    "stockOh": 0,
                    "stockOnOrder": 0,
                }
            ],
        }
    ]
    content = build_performance_workbook(
        report,
        mode="sales",
        metric="amount",
        grain="day_total",
        show_descriptions=False,
        branch_ids=["60016", "60926"],
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Report"]

    assert sheet["F5"].value == "17/08/2026"
    assert sheet["F6"].value == 150.5
    assert sheet["E6"].value == "=SUM(F6:F6)"
    workbook.close()

def test_performance_export_filename_has_view_and_unique_timestamp() -> None:
    filename = performance_export_filename(
        "sales",
        "amount",
        "day",
        datetime(2026, 8, 25, 12, 34, 56),
    )

    assert filename == "TWD_Sales_Amount_Branch_20260825_123456.xlsx"

    assert performance_export_filename(
        "sales",
        "amount",
        "day",
        datetime(2026, 8, 29, 17, 30, tzinfo=UTC),
    ) == "TWD_Sales_Amount_Branch_20260830_003000.xlsx"


def test_gross_sales_export_is_clearly_labeled() -> None:
    content = build_performance_workbook(
        _report(),
        mode="sales",
        metric="amount",
        grain="day",
        show_descriptions=True,
        sales_basis="gross",
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    assert workbook["Report"]["A1"].value == (
        "TWD Sales (Gross Sale Out) by Branch — Amount"
    )
    workbook.close()

    assert performance_export_filename(
        "sales",
        "amount",
        "day",
        datetime(2026, 8, 25, 12, 34, 56),
        sales_basis="gross",
    ) == "TWD_Sales_Gross_Amount_Branch_20260825_123456.xlsx"


@pytest.mark.parametrize("mt_code", ["TWD", "HP", "MH"])
def test_inventory_export_contains_turnover_and_exact_selected_ranges(mt_code: str) -> None:
    report = _report()
    report["selectedMonth"] = None
    report["selectedDateRanges"] = [
        {"from": "2026-06-01", "to": "2026-06-30"},
        {"from": "2026-08-17", "to": "2026-08-17"},
    ]
    report["inventorySummary"] = {"averageTom": 5.99, "averageTod": 179.7}
    report["items"] = [{**report["items"][0], "tom": 5.99, "tod": 179.7}]

    content = build_performance_workbook(
        report,
        mode="inventory",
        metric="stockOh",
        grain="day",
        show_descriptions=True,
        mt_code=mt_code,
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Report"]

    assert [sheet.cell(5, column).value for column in range(1, 10)] == [
        "Sho",
        "Pro",
        f"{mt_code} SKU",
        f"{mt_code} Description",
        "WA Item",
        "WA Description",
        "TOM",
        "TOD",
        "TOTAL",
    ]
    assert sheet["A4"].value == "AVG"
    assert sheet["G4"].value == 5.99
    assert sheet["H4"].value == 179.7
    assert sheet["G6"].value == 5.99
    assert sheet["H6"].value == 179.7
    assert sheet.freeze_panes == "J6"
    assert "01/06/2026 – 30/06/2026, 17/08/2026" in sheet["A2"].value
    workbook.close()


def test_inventory_month_export_matches_monthly_snapshot_points() -> None:
    report = _report()
    report["dates"] = ["2026-07", "2026-08"]
    report["selectedMonth"] = None
    report["inventorySummary"] = {"averageTom": 5.99, "averageTod": 179.7}
    report["items"] = [
        {
            **report["items"][0],
            "tom": 5.99,
            "tod": 179.7,
            "points": [
                {
                    "date": "2026-07",
                    "branchId": "all",
                    "amount": 0,
                    "qty": 0,
                    "stockOh": 10,
                    "stockOnOrder": 3,
                },
                {
                    "date": "2026-08",
                    "branchId": "all",
                    "amount": 0,
                    "qty": 0,
                    "stockOh": 20,
                    "stockOnOrder": 7,
                },
            ],
        }
    ]

    content = build_performance_workbook(
        report,
        mode="inventory",
        metric="stockOh",
        grain="month",
        show_descriptions=True,
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Report"]

    assert sheet["J5"].value == "2026-07"
    assert sheet["K5"].value == "2026-08"
    assert sheet["J6"].value == 10
    assert sheet["K6"].value == 20
    assert sheet["I6"].value == "=SUM(J6:K6)"
    assert "Jul 2026 – Aug 2026" in sheet["A2"].value
    workbook.close()
