from collections import Counter
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
import xlrd
from openpyxl import Workbook

from app.importers.twd import (
    TwdFormatError,
    _decimal_cell,
    _source_cell_warnings,
    calculate_amount,
    extract_twd_file,
)
from app.models import ImportBatch

SAMPLE_DIR = Path(__file__).parents[2] / ".tmp" / "twd-samples"


def test_calculate_amount_uses_decimal() -> None:
    assert calculate_amount(Decimal("107")) == Decimal("100.000000000000")


def test_import_status_values_fit_database_column() -> None:
    status_length = ImportBatch.__table__.c.status.type.length
    assert status_length is not None
    assert (
        max(map(len, ("imported", "imported_with_warnings", "failed", "duplicate")))
        <= status_length
    )


def test_excel_error_cell_is_not_treated_as_a_quantity() -> None:
    errors: Counter[tuple[str, str]] = Counter()
    cell = SimpleNamespace(ctype=xlrd.XL_CELL_ERROR, value=15)

    assert _decimal_cell(cell, "Stock On Hand", errors) == Decimal("0")
    assert _source_cell_warnings(errors) == (
        "Stock On Hand: พบเซลล์ต้นทาง #VALUE! จำนวน 1 เซลล์ ระบบไม่นำมารวมยอด",
    )


def test_xlsx_uses_same_twd_contract_and_reads_period_from_file(tmp_path) -> None:
    path = tmp_path / "folder-date-does-not-matter.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ReportSaleSubscription"
    sheet.cell(row=3, column=2, value="Mon,31 Aug 2026")
    headers = {
        1: "Store",
        3: "Cat",
        4: "Sub Cat",
        5: "Brand",
        6: "SKU",
        7: "Barcode",
        8: "Description",
        10: "Product Type",
        11: "Sales Amount",
        12: "Sales Qty",
        13: "Stock On Hand",
        14: "Stock On Order",
        15: "Last Sold Date",
        16: "Last Receive Date",
    }
    for column, value in headers.items():
        sheet.cell(row=6, column=column, value=value)
    values = {
        1: "60016-ภูเก็ต เฟสติวัล",
        3: "001",
        4: "002",
        5: "Brand",
        6: "60365148",
        7: "8850000000000",
        8: "สินค้าทดลอง",
        10: "Product",
        11: 107,
        12: 1,
        13: "#VALUE!",
        14: 2,
        15: "2026-08-30",
        16: "2026-08-29",
    }
    for column, value in values.items():
        sheet.cell(row=7, column=column, value=value)
    sheet.cell(row=8, column=8, value="Total")
    sheet.cell(row=8, column=11, value=107)
    sheet.cell(row=8, column=12, value=1)
    sheet.cell(row=8, column=13, value=0)
    sheet.cell(row=8, column=14, value=2)
    sheet.cell(row=9, column=1, value="Count of Rows - 1")
    workbook.save(path)

    extract = extract_twd_file(path)

    assert extract.data_date.isoformat() == "2026-08-31"
    assert extract.summary.row_count == 1
    assert extract.summary.amount == Decimal("100.000000000000")
    assert extract.summary.stock_on_hand == Decimal("0")
    assert extract.reconciliation_errors == (
        "Stock On Hand: พบเซลล์ต้นทาง #VALUE! จำนวน 1 เซลล์ ระบบไม่นำมารวมยอด",
    )


def test_empty_xls_is_reported_as_file_format_error(tmp_path) -> None:
    path = tmp_path / "empty.xls"
    path.write_bytes(b"")

    with pytest.raises(
        TwdFormatError,
        match=r"เปิดไฟล์ empty\.xls ไม่สำเร็จ: File size is 0 bytes",
    ):
        extract_twd_file(path)


@pytest.mark.parametrize(
    (
        "filename",
        "period",
        "rows",
        "source_amount",
        "qty",
        "stock_oh",
        "reported_stock_oh",
        "stock_order",
        "negative",
    ),
    [
        (
            "twd-2026-08-16.xls",
            "2026-08-16",
            12560,
            "1101053.72",
            "353",
            "77265",
            "77265",
            "5902",
            4,
        ),
        (
            "twd-2026-08-17.xls",
            "2026-08-17",
            12561,
            "864345.32",
            "345",
            "77416",
            "77416",
            "6338",
            2,
        ),
    ],
)
def test_real_twd_samples_reconcile(
    filename, period, rows, source_amount, qty, stock_oh, reported_stock_oh, stock_order, negative
) -> None:
    path = SAMPLE_DIR / filename
    if not path.exists():
        pytest.skip("ไฟล์ตัวอย่างจาก NAS ไม่มีใน workspace")
    extract = extract_twd_file(path)
    summary = extract.summary
    assert extract.data_date.isoformat() == period
    assert (summary.row_count, summary.store_count, summary.sku_count) == (rows, 102, 2043)
    assert summary.source_amount == Decimal(source_amount)
    assert summary.sales_qty == Decimal(qty)
    assert summary.stock_on_hand == Decimal(stock_oh)
    assert extract.reported_summary.stock_on_hand == Decimal(reported_stock_oh)
    assert summary.stock_on_order == Decimal(stock_order)
    assert summary.negative_row_count == negative
    assert extract.reconciliation_errors == (
        "Stock On Hand: พบเซลล์ต้นทาง #VALUE! จำนวน 35 เซลล์ ระบบไม่นำมารวมยอด",
    )
    assert extract.rows[0].description
    assert any("\u0e00" <= character <= "\u0e7f" for character in extract.rows[0].description)
